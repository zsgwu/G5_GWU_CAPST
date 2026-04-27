import csv
import difflib
import hashlib
import html
import json
import logging
import os
import re
import sqlite3
from io import BytesIO
from pathlib import Path

import altair as alt
import openai
import pandas as pd
import pdfplumber
import streamlit as st
from docx import Document
from rag_retrival import looks_like_rag_question, retrieve_context as retrieve_rag_context

BASE_SYSTEM_MESSAGE = (
    "You are an intelligent orchestration agent for an education advisory chatbot. "
    "Your job is to understand the user's intent, select the most appropriate data source or sources, "
    "and generate the best possible response. "
    "Available sources: optional resume context for personalized advice, career guidance, and application strategy; "
    "optional applicant stats such as GPA or SAT/ACT for admissions chances, school recommendations, and fit analysis; "
    "retrieved RAG program context for majors, career outcomes, salaries, job placement, and program-level insights; "
    "SQLite school data for acceptance rate, tuition, enrollment, and other school-level institutional statistics only; "
    "and general Azure reasoning when no structured source is relevant. "
    "Silently classify each prompt as PERSONALIZED_ADVICE, MAJOR_OR_CAREER, SCHOOL_STATS, or GENERAL. "
    "Follow this priority: personalized advice first, then major or career questions, then school stats, then general reasoning. "
    "Do not default to SQL, do not force structured data when unnecessary, do not block the answer if resume or stats are missing, "
    "and if unsure do not use SQL. "
    "Do not use SQL for major-specific questions, career-outcome questions, vague questions, or pure advice questions unless "
    "school-level context is explicitly being used as supporting context. "
    "If extra information would help, offer it as a low-pressure option rather than a requirement."
)
AZURE_RESPONSE_STYLE_SYSTEM_MESSAGE = (
    "Before answering, do a quick silent check that the answer is helpful for a student decision, "
    "directly answers the question, and is as concise as possible. "
    "Then give the final answer only. Do not show the self-check. "
    "Keep responses concise and decision-oriented. Default to 1-3 short paragraphs. "
    "Only use bullets when the user explicitly asks for a comparison, list, or options, and keep those to 2-4 flat bullets maximum. "
    "Avoid filler, repetition, long preambles, and generic encouragement. If there is uncertainty, say it plainly. "
    "Finish the answer cleanly and do not stop mid-thought."
)
SQLITE_CONTEXT_SYSTEM_MESSAGE = (
    "Use the SQLite-backed College Scorecard context below as the primary source for numeric claims. "
    "This database is institution-level only. Do not claim program-level or major-level outcomes unless "
    "the context explicitly includes them. If the user asks for a field that is not present, say so clearly."
)
RAG_CONTEXT_SYSTEM_MESSAGE = (
    "Use the retrieved program context below only for claims about majors, CIP families, degree levels, "
    "related occupations, and national program earnings. Do not turn retrieved program context into "
    "school-specific claims unless the context explicitly says that."
)

PII_PATTERNS = {
    "email": re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"),
    "phone": re.compile(r"\b(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}\b"),
    "linkedin": re.compile(
        r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s,]+", re.IGNORECASE
    ),
    "address": re.compile(
        r"\b\d{1,5}\s+[A-Za-z0-9\.\-'\s]{2,}\b(?:Street|St|Road|Rd|Avenue|Ave|Boulevard|Blvd|"
        r"Lane|Ln|Drive|Dr|Court|Ct|Circle|Plaza|Terrace|Terr|Parkway|Pkwy|Trail|Trl|Highway|Hwy|"
        r"Route|Rte)\b[^\n]*",
        re.IGNORECASE,
    ),
}

LOCATION_LINE_PATTERN = re.compile(r"(?m)(^|\n)([A-Z][A-Za-z .'-]+,\s*[A-Z]{2}\b)")
NULL_TOKENS = {"", "null", "privacysuppressed", "ps", "na", "n/a"}

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

DATASET_PATH = Path(__file__).with_name("Most-Recent-Cohorts-Institution.csv")
SUPPLEMENTAL_NET_PRICE_PATH = Path(__file__).with_name("instate_tuition_cleaned.csv")
SQLITE_PATH = Path(__file__).with_name("most_recent_cohorts_institution.sqlite3")
LOGO_PATH = Path(__file__).with_name("EDyoU.jpg")
SCORECARD_TABLE = "scorecard_institutions"
SUPPLEMENTAL_NET_PRICE_TABLE = "supplemental_net_price_history"
SQLITE_IMPORT_BATCH_SIZE = 1000
DEFAULT_METRIC_KEY = "earnings"
DEFAULT_RESULT_LIMIT = 8
MAX_RESULT_LIMIT = 20
MAX_HISTORY_TURNS_FOR_MODEL = 6
AZURE_MAX_RESPONSE_TOKENS = 700
AZURE_REASONING_EFFORT = "low"
MAX_SQL_SUMMARY_WORDS = 55
MAX_SQL_SUMMARY_SENTENCES = 2
EXCLUDED_SQL_UNITIDS = {
    178721,  # Park University (Parkville, MO)
}
SUPPLEMENTAL_NET_PRICE_VIEW_OPTIONS = (
    ("average_net_price", "Average net price"),
    ("net_price_income_0_30000", "$0-30k"),
    ("net_price_income_30001_48000", "$30,001-48k"),
    ("net_price_income_48001_75000", "$48,001-75k"),
    ("net_price_income_75001_110000", "$75,001-110k"),
    ("net_price_income_110001_plus", "$110,001+"),
)
SUPPLEMENTAL_NET_PRICE_VIEW_LABELS = dict(SUPPLEMENTAL_NET_PRICE_VIEW_OPTIONS)
SUPPLEMENTAL_NET_PRICE_DEFAULT_VIEW = "average_net_price"
SUPPLEMENTAL_NET_PRICE_CHART_YEARS = 5

SELECTED_COLUMNS = [
    "UNITID",
    "INSTNM",
    "CITY",
    "STABBR",
    "CONTROL",
    "LOCALE",
    "CCBASIC",
    "PREDDEG",
    "HIGHDEG",
    "ADM_RATE",
    "UGDS",
    "TUITIONFEE_IN",
    "TUITIONFEE_OUT",
    "PCTPELL",
    "PCTFLOAN",
    "DEBT_MDN",
    "OMAWDP8_ALL",
    "MD_EARN_WNE_P10",
]

MATCH_STOPWORDS = {
    "a",
    "about",
    "admission",
    "admissions",
    "after",
    "aid",
    "all",
    "and",
    "are",
    "at",
    "best",
    "by",
    "college",
    "colleges",
    "compare",
    "comparison",
    "cost",
    "debt",
    "for",
    "from",
    "get",
    "graph",
    "highest",
    "how",
    "in",
    "institution",
    "institutions",
    "is",
    "it",
    "largest",
    "lowest",
    "me",
    "most",
    "net",
    "of",
    "on",
    "price",
    "rate",
    "school",
    "schools",
    "show",
    "the",
    "their",
    "these",
    "this",
    "to",
    "top",
    "tuition",
    "university",
    "universities",
    "vs",
    "what",
    "which",
    "with",
}

CONTROL_LABELS = {
    "1": "Public",
    "2": "Private nonprofit",
    "3": "Private for-profit",
}
CONTROL_SCOPE_LABELS = {
    ("1",): "Public",
    ("2",): "Private nonprofit",
    ("3",): "Private for-profit",
    ("2", "3"): "Private",
}

STATE_NAME_TO_ABBR = {
    "alabama": "AL",
    "alaska": "AK",
    "arizona": "AZ",
    "arkansas": "AR",
    "california": "CA",
    "colorado": "CO",
    "connecticut": "CT",
    "delaware": "DE",
    "district of columbia": "DC",
    "american samoa": "AS",
    "federated states of micronesia": "FM",
    "florida": "FL",
    "georgia": "GA",
    "guam": "GU",
    "hawaii": "HI",
    "idaho": "ID",
    "illinois": "IL",
    "indiana": "IN",
    "iowa": "IA",
    "kansas": "KS",
    "kentucky": "KY",
    "louisiana": "LA",
    "maine": "ME",
    "maryland": "MD",
    "massachusetts": "MA",
    "michigan": "MI",
    "minnesota": "MN",
    "mississippi": "MS",
    "missouri": "MO",
    "montana": "MT",
    "marshall islands": "MH",
    "nebraska": "NE",
    "nevada": "NV",
    "new hampshire": "NH",
    "new jersey": "NJ",
    "new mexico": "NM",
    "new york": "NY",
    "northern mariana islands": "MP",
    "north carolina": "NC",
    "north dakota": "ND",
    "ohio": "OH",
    "oklahoma": "OK",
    "oregon": "OR",
    "palau": "PW",
    "pennsylvania": "PA",
    "puerto rico": "PR",
    "rhode island": "RI",
    "south carolina": "SC",
    "south dakota": "SD",
    "tennessee": "TN",
    "texas": "TX",
    "utah": "UT",
    "vermont": "VT",
    "virginia": "VA",
    "washington": "WA",
    "west virginia": "WV",
    "wisconsin": "WI",
    "virgin islands": "VI",
    "u.s. virgin islands": "VI",
    "us virgin islands": "VI",
    "wyoming": "WY",
}
STATE_ABBRS = set(STATE_NAME_TO_ABBR.values())
STATE_ABBR_TO_NAME: dict[str, str] = {}
for state_name, abbreviation in STATE_NAME_TO_ABBR.items():
    STATE_ABBR_TO_NAME.setdefault(abbreviation, state_name.title())
STATE_ABBR_TO_NAME["DC"] = "District of Columbia"
STATE_ABBR_TO_NAME["VI"] = "U.S. Virgin Islands"
DMV_STATE_ABBRS = ("DC", "MD", "VA")
SCHOOL_NAME_FILLER_TOKENS = {
    "the",
    "of",
    "at",
    "and",
    "for",
    "campus",
    "branch",
    "main",
    "system",
}
INSTITUTION_TYPE_TOKENS = {"university", "college", "institute", "school"}
FOLLOW_UP_SCHOOL_PHRASES = (
    "what about",
    "how about",
    "what is its",
    "what's its",
    "what is their",
    "what's their",
    "how big is it",
    "how big are they",
    "how large is it",
    "how large are they",
    "compare it",
    "compare them",
    "against them",
    "against it",
    "same school",
    "same schools",
    "those schools",
    "these schools",
)
DIRECT_CONTEXT_REFERENCE_PHRASES = (
    "their",
    "them",
    "those schools",
    "these schools",
    "same school",
    "same schools",
    "it",
    "its",
)
DATA_PROFILE_INTENT_PHRASES = (
    "tell me about",
    "show me",
    "show data for",
    "show stats for",
    "give me data on",
    "give me stats for",
    "overview of",
    "profile of",
    "school profile",
    "school overview",
)
DATA_COMPARISON_INTENT_PHRASES = (
    "compare",
    "versus",
    " vs ",
    "against",
    "relative to",
)
MODEL_PREFERRED_INTENT_PHRASES = (
    "better fit",
    "fit for me",
    "fit with my",
    "should i",
    "which one should",
    "which school should",
    "what do you think",
    "recommend",
    "student life",
    "campus life",
    "social life",
    "campus culture",
    "culture",
    "vibe",
    "feel like",
    "essay",
    "essays",
)
PERSONALIZED_ADVICE_INTENT_PHRASES = (
    "would i get in",
    "could i get in",
    "can i get into",
    "can i get in",
    "my chances",
    "chance me",
    "should i apply",
    "should i go",
    "what schools should i",
    "which schools should i",
    "which school is better for me",
    "which is better for me",
    "best fit for me",
    "good fit for me",
    "fit me best",
    "based on my",
    "with my gpa",
    "with my sat",
    "with my act",
    "for someone like me",
    "for my profile",
    "my profile",
)
FIRST_PERSON_REFERENCE_PATTERN = re.compile(r"\b(i|me|my|mine|myself)\b", re.IGNORECASE)
APPLICANT_STAT_QUERY_TOKENS = (
    "gpa",
    "sat",
    "act",
    "my stats",
    "my score",
    "my scores",
)
INTENT_PERSONALIZED_ADVICE = "PERSONALIZED_ADVICE"
INTENT_MAJOR_OR_CAREER = "MAJOR_OR_CAREER"
INTENT_SCHOOL_STATS = "SCHOOL_STATS"
INTENT_GENERAL = "GENERAL"
BLOCKED_GEOGRAPHIC_ALIASES = set(STATE_NAME_TO_ABBR.keys()) | {
    "united states",
    "america",
}
END_SESSION_NEXT_STEP_VERBS = (
    "Build",
    "Compare",
    "Explore",
    "Research",
    "Review",
    "Refine",
    "Check",
    "Ask",
    "Visit",
    "Reach",
    "Contact",
    "Map",
    "Shortlist",
    "Look",
)
END_SESSION_BLOCKED_PHRASES = (
    "chat history",
    "session",
    "transcript",
    "database",
    "sql",
    "sqlite",
    "repository",
    "query",
)
END_SESSION_SUMMARY_SYSTEM_MESSAGE = (
    "You are an education advisor providing a thoughtful end-of-session summary for a student. "
    "Your goals are to synthesize the conversation into meaningful insights, highlight patterns around fit, direction, "
    "and decision-making, and provide clear, motivating next steps. "
    "You will receive the conversation transcript plus optional resume context and optional applicant stats. "
    "The user may not have provided resume or stats, and that is completely fine. "
    "Return strict JSON only with keys `top_takeaways`, `next_steps`, and `closing_thought`. "
    "`top_takeaways` must be an array of exactly 3 concise strings. "
    "`next_steps` must be an array of exactly 3 concise strings. "
    "`closing_thought` must be one short natural sentence. "
    "For top takeaways: focus on fit, direction, and positioning; reflect what the student is learning about themselves; "
    "be forward-looking and thoughtful; do not list raw stats; do not repeat database outputs; avoid generic fluff. "
    "For next steps: make them concrete, actionable, realistic, and not overwhelming. "
    "Before finalizing each next step, silently ask whether it is actionable and realistic, then rewrite anything vague. "
    "Do not list specific queried statistics. Do not restate the entire conversation. "
    "Do not mention the chat history, transcript, session, database, SQL, repository, or missing data. "
    "If resume or applicant stats were clearly useful, lightly incorporate them into the interpretation, but keep the focus on meaning rather than numbers."
)

NET_PRICE_SQL = f"""
(
    SELECT AVG(s.average_net_price)
    FROM {SUPPLEMENTAL_NET_PRICE_TABLE} AS s
    WHERE s.normalized_name = {SCORECARD_TABLE}.NORMALIZED_INSTNM
      AND s.year_start = (
          SELECT MAX(s2.year_start)
          FROM {SUPPLEMENTAL_NET_PRICE_TABLE} AS s2
          WHERE s2.normalized_name = {SCORECARD_TABLE}.NORMALIZED_INSTNM
      )
)
"""

COMMON_SELECT_SQL = f"""
SELECT
    UNITID,
    INSTNM,
    NORMALIZED_INSTNM,
    CITY,
    STABBR,
    CONTROL,
    LOCALE,
    CCBASIC,
    PREDDEG,
    HIGHDEG,
    ADM_RATE,
    UGDS,
    TUITIONFEE_IN,
    TUITIONFEE_OUT,
    PCTPELL,
    PCTFLOAN,
    DEBT_MDN,
    OMAWDP8_ALL,
    MD_EARN_WNE_P10,
    {NET_PRICE_SQL} AS net_price
FROM {SCORECARD_TABLE}
"""

METRIC_DEFINITIONS = {
    "earnings": {
        "label": "Median Earnings (10 Years After Entry)",
        "short_label": "Median earnings",
        "expression": "MD_EARN_WNE_P10",
        "format": "currency",
        "default_direction": "DESC",
    },
    "net_price": {
        "label": "Average Net Price Most Recent Year",
        "short_label": "Average net price",
        "expression": NET_PRICE_SQL,
        "format": "currency",
        "default_direction": "ASC",
    },
    "tuition_in": {
        "label": "In-State Tuition and Fees",
        "short_label": "In-state tuition",
        "expression": "TUITIONFEE_IN",
        "format": "currency",
        "default_direction": "ASC",
    },
    "tuition_out": {
        "label": "Out-of-State Tuition and Fees",
        "short_label": "Out-of-state tuition",
        "expression": "COALESCE(TUITIONFEE_OUT, TUITIONFEE_IN)",
        "format": "currency",
        "default_direction": "ASC",
    },
    "debt": {
        "label": "Median Federal Debt",
        "short_label": "Median debt",
        "expression": "DEBT_MDN",
        "format": "currency",
        "default_direction": "ASC",
    },
    "admission_rate": {
        "label": "Admission Rate",
        "short_label": "Admission rate",
        "expression": "ADM_RATE",
        "format": "percent",
        "default_direction": "DESC",
    },
    "completion_rate": {
        "label": "Eight-Year Completion / Award Rate",
        "short_label": "Completion rate",
        "expression": "OMAWDP8_ALL",
        "format": "percent",
        "default_direction": "DESC",
    },
    "pell_share": {
        "label": "Pell Grant Share",
        "short_label": "Pell share",
        "expression": "PCTPELL",
        "format": "percent",
        "default_direction": "DESC",
    },
    "loan_share": {
        "label": "Federal Loan Share",
        "short_label": "Loan share",
        "expression": "PCTFLOAN",
        "format": "percent",
        "default_direction": "DESC",
    },
    "undergrad_size": {
        "label": "Undergraduate Enrollment",
        "short_label": "Undergraduates",
        "expression": "UGDS",
        "format": "count",
        "default_direction": "DESC",
    },
}

st.set_page_config(
    page_title="Decision Intelligence Chatbot",
    page_icon="EDyoU.jpg",
    layout="wide",
)


def _apply_brand_theme() -> None:
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=League+Spartan:wght@500;600;700&family=DM+Sans:wght@400;500;700&display=swap');

        :root {
            --ey-bg: #ebe6da;
            --ey-bg-soft: #f5f0e6;
            --ey-surface: rgba(255, 251, 245, 0.76);
            --ey-surface-strong: rgba(255, 251, 245, 0.92);
            --ey-line: rgba(16, 82, 62, 0.12);
            --ey-line-strong: rgba(16, 82, 62, 0.22);
            --ey-ink: #16382e;
            --ey-muted: #5d6a64;
            --ey-accent: #10523e;
            --ey-accent-soft: #8cb885;
            --ey-sand: #c8c1ac;
        }

        html, body, [class*="stApp"] {
            font-family: "DM Sans", sans-serif;
            background:
                radial-gradient(circle at top left, rgba(140, 184, 133, 0.18), transparent 28%),
                linear-gradient(180deg, #f7f3ea 0%, var(--ey-bg) 42%, #e4ddcf 100%);
            color: var(--ey-ink);
            scroll-padding-top: 6rem;
        }

        .stApp {
            color: var(--ey-ink);
        }

        .block-container {
            padding-top: 1.3rem;
            padding-bottom: 7rem;
            max-width: 1600px;
        }

        section[data-testid="stSidebar"] {
            width: min(25vw, 390px) !important;
            min-width: 300px !important;
            border-right: 1px solid var(--ey-line);
            background:
                linear-gradient(180deg, rgba(255, 251, 245, 0.96) 0%, rgba(244, 239, 228, 0.98) 100%);
        }

        section[data-testid="stSidebar"] > div {
            background: transparent;
        }

        [data-testid="stSidebarUserContent"] {
            padding-top: 1.2rem;
        }

        h1, h2, h3, h4 {
            font-family: "League Spartan", sans-serif;
            color: var(--ey-ink);
            letter-spacing: -0.03em;
        }

        p, li, label, .stCaption, .stMarkdown, .stText, .stTextArea textarea {
            font-family: "DM Sans", sans-serif;
        }

        [data-testid="stTabs"] {
            gap: 0.6rem;
        }

        [data-testid="stTabs"] button[role="tab"] {
            border-radius: 999px;
            border: 1px solid var(--ey-line);
            background: rgba(255, 251, 245, 0.55);
            color: var(--ey-muted);
            padding: 0.5rem 1rem;
            font-weight: 600;
        }

        [data-testid="stTabs"] button[role="tab"][aria-selected="true"] {
            background: var(--ey-accent);
            color: white;
            border-color: var(--ey-accent);
        }

        [data-testid="stPills"] [role="radiogroup"] {
            gap: 0.38rem;
            flex-wrap: wrap;
        }

        [data-testid="stPills"] button {
            min-height: 1.85rem !important;
            padding: 0.16rem 0.72rem !important;
            border-radius: 999px !important;
            border: 1px solid var(--ey-line) !important;
            background: rgba(255, 251, 245, 0.72) !important;
            color: var(--ey-muted) !important;
            font-weight: 600 !important;
            box-shadow: none !important;
        }

        [data-testid="stPills"] button[aria-pressed="true"],
        [data-testid="stPills"] button[aria-checked="true"] {
            background: rgba(16, 82, 62, 0.08) !important;
            color: var(--ey-ink) !important;
            border-color: rgba(16, 82, 62, 0.18) !important;
        }

        [data-testid="stChatMessage"] {
            background: var(--ey-surface);
            border: 1px solid var(--ey-line);
            border-radius: 22px;
            padding: 0.8rem 1rem;
            margin-bottom: 0.8rem;
            backdrop-filter: blur(14px);
            scroll-margin-top: 6rem;
        }

        [data-testid="stChatMessageContent"] {
            width: 100%;
            min-width: 0;
        }

        [data-testid="stChatMessageContent"] > div,
        [data-testid="stChatMessageContent"] .stMarkdown,
        [data-testid="stChatMessageContent"] .stMarkdownContainer {
            width: 100%;
            min-width: 0;
            max-width: 100%;
        }

        [data-testid="stChatMessageContent"] h1,
        [data-testid="stChatMessageContent"] h2,
        [data-testid="stChatMessageContent"] h3,
        [data-testid="stChatMessageContent"] h4,
        [data-testid="stChatMessageContent"] h5,
        [data-testid="stChatMessageContent"] h6 {
            font-family: "DM Sans", sans-serif !important;
            font-size: 1rem !important;
            line-height: 1.45 !important;
            letter-spacing: 0 !important;
            margin: 0.2rem 0 0.55rem !important;
        }

        [data-testid="stChatMessageContent"] p,
        [data-testid="stChatMessageContent"] li,
        [data-testid="stChatMessageContent"] blockquote,
        [data-testid="stChatMessageContent"] label,
        [data-testid="stChatMessageContent"] span,
        [data-testid="stChatMessageContent"] div {
            overflow-wrap: anywhere;
            word-break: break-word;
        }

        [data-testid="stChatMessageContent"] pre,
        [data-testid="stChatMessageContent"] code,
        [data-testid="stChatMessageContent"] table {
            max-width: 100%;
        }

        [data-testid="stChatMessageContent"] pre {
            white-space: pre-wrap;
            overflow-x: auto;
        }

        [data-testid="stChatMessageContent"] table {
            display: block;
            overflow-x: auto;
        }

        [data-testid="stMetric"] {
            background: var(--ey-surface);
            border: 1px solid var(--ey-line);
            border-radius: 20px;
            padding: 0.6rem 0.8rem;
            backdrop-filter: blur(14px);
        }

        [data-testid="stDataFrame"] {
            border-radius: 20px;
            overflow: hidden;
            border: 1px solid var(--ey-line);
        }

        [data-testid="stButton"] > button {
            border-radius: 999px;
            background: var(--ey-accent);
            color: white;
            border: 1px solid var(--ey-accent);
            font-weight: 700;
            min-height: 2.85rem;
        }

        [data-testid="stButton"] > button:hover {
            background: #0c4534;
            border-color: #0c4534;
        }

        [data-testid="stChatInput"] {
            background: linear-gradient(180deg, rgba(235, 230, 218, 0), rgba(235, 230, 218, 0.95) 28%, rgba(235, 230, 218, 1) 100%);
            padding-top: 1rem;
        }

        [data-testid="stChatInput"] textarea {
            background: var(--ey-surface-strong);
            border: 1px solid var(--ey-line-strong);
            border-radius: 22px;
            color: var(--ey-ink);
            box-shadow: 0 12px 34px rgba(16, 82, 62, 0.08);
        }

        [data-testid="stExpander"] {
            border-radius: 18px;
            border: 1px solid var(--ey-line);
            background: rgba(255, 251, 245, 0.44);
        }

        .edu-shell {
            background: rgba(255, 251, 245, 0.58);
            border: 1px solid var(--ey-line);
            border-radius: 26px;
            padding: 0.95rem 1.15rem 0.9rem;
            backdrop-filter: blur(12px);
            box-shadow: 0 18px 40px rgba(16, 82, 62, 0.08);
        }

        .edu-brand-row {
            display: flex;
            align-items: center;
            gap: 1rem;
            margin-bottom: 1rem;
        }

        .edu-kicker {
            text-transform: uppercase;
            letter-spacing: 0.18em;
            font-size: 0.72rem;
            font-weight: 700;
            color: var(--ey-accent);
            margin-bottom: 0.3rem;
        }

        .edu-title {
            font-family: "League Spartan", sans-serif;
            font-size: clamp(1.6rem, 2.4vw, 2.45rem);
            line-height: 1.02;
            color: var(--ey-ink);
            margin: 0;
        }

        .edu-subtitle {
            margin-top: 0.3rem;
            color: var(--ey-muted);
            max-width: 42rem;
            font-size: 0.96rem;
            line-height: 1.42;
        }

        .edu-rail {
            background: rgba(255, 251, 245, 0.52);
            border: 1px solid var(--ey-line);
            border-radius: 28px;
            padding: 0.95rem 1rem;
            backdrop-filter: blur(12px);
        }

        .edu-scroll-note {
            color: var(--ey-muted);
            font-size: 0.9rem;
            margin-bottom: 0.55rem;
        }

        .edu-chat-top-guard {
            height: 4.5rem;
            margin: 0 0 1rem;
            border-radius: 26px;
            border: 1px solid rgba(16, 82, 62, 0.08);
            background:
                linear-gradient(180deg, rgba(255, 251, 245, 0.9) 0%, rgba(255, 251, 245, 0.55) 58%, rgba(255, 251, 245, 0) 100%);
            pointer-events: none;
        }

        .edu-summary-stat {
            margin-bottom: 1rem;
        }

        .edu-summary-label {
            color: var(--ey-muted);
            font-size: 0.82rem;
            font-weight: 600;
            letter-spacing: 0.02em;
            margin-bottom: 0.18rem;
        }

        .edu-summary-value {
            font-family: "League Spartan", sans-serif;
            font-size: clamp(1.65rem, 2.1vw, 2.4rem);
            line-height: 1.02;
            color: var(--ey-ink);
        }

        .edu-result-summary {
            margin: 0.15rem 0 1rem;
            max-width: 58rem;
            color: var(--ey-ink);
            font-family: "DM Sans", sans-serif;
            font-size: 0.98rem;
            line-height: 1.58;
            white-space: pre-wrap;
            overflow-wrap: anywhere;
            word-break: break-word;
        }

        .edu-rag-watermark {
            margin: 0.65rem 0 0.15rem;
            padding-top: 0.55rem;
            border-top: 1px solid rgba(16, 82, 62, 0.12);
            color: var(--ey-muted);
            font-family: "DM Sans", sans-serif;
            font-size: 0.76rem;
            line-height: 1.45;
            overflow-wrap: anywhere;
            word-break: break-word;
        }

        .edu-rag-watermark strong {
            color: var(--ey-ink);
            font-weight: 700;
        }

        .edu-chat-copy {
            color: var(--ey-ink);
            font-family: "DM Sans", sans-serif;
            font-size: 0.98rem;
            line-height: 1.58;
            white-space: pre-wrap;
            overflow-wrap: anywhere;
            word-break: break-word;
        }

        .edu-sidebar-logo {
            width: 100%;
            margin: 0 0 1.15rem;
        }

        .edu-sidebar-logo [data-testid="stImage"] {
            width: 100%;
        }

        .edu-sidebar-logo img {
            width: 100% !important;
            height: auto;
            display: block;
        }

        .edu-sidebar-title {
            font-family: "League Spartan", sans-serif;
            font-size: 1.55rem;
            line-height: 1;
            color: var(--ey-ink);
            margin-bottom: 0.85rem;
        }

        .edu-sidebar-section-title {
            font-family: "League Spartan", sans-serif;
            font-size: 1.15rem;
            line-height: 1.05;
            color: var(--ey-ink);
            margin: 0 0 0.35rem;
        }

        .edu-sidebar-section-copy {
            color: var(--ey-muted);
            font-size: 0.82rem;
            line-height: 1.45;
            margin: 0 0 0.95rem;
        }

        .edu-sidebar-field-group {
            margin-bottom: 1rem;
        }

        .edu-sidebar-nav-label {
            text-transform: uppercase;
            letter-spacing: 0.16em;
            color: var(--ey-muted);
            font-size: 0.7rem;
            font-weight: 700;
            margin: 1rem 0 0.55rem;
        }

        .edu-sidebar-resume {
            margin-bottom: 1rem;
        }

        section[data-testid="stSidebar"] [data-testid="stExpander"] {
            margin-bottom: 0.9rem;
        }

        .edu-main-shell {
            min-height: 100%;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def _render_sidebar_panel() -> bool:
    with st.sidebar:
        if LOGO_PATH.exists():
            st.markdown('<div class="edu-sidebar-logo">', unsafe_allow_html=True)
            st.image(str(LOGO_PATH), use_container_width=True)
            st.markdown("</div>", unsafe_allow_html=True)
        _render_application_information_panel()
        st.markdown('<div class="edu-sidebar-resume">', unsafe_allow_html=True)
        _render_resume_panel()
        st.markdown("</div>", unsafe_allow_html=True)
        end_session_clicked = st.button("End Session", use_container_width=True)
        st.divider()
        return end_session_clicked


def _render_application_information_panel() -> None:
    st.markdown(
        '<div class="edu-sidebar-section-title">Application Information</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        (
            '<div class="edu-sidebar-section-copy">'
            "Please input relevant information here to improve accuracy of chatbot experience."
            "</div>"
        ),
        unsafe_allow_html=True,
    )
    st.markdown('<div class="edu-sidebar-field-group">', unsafe_allow_html=True)
    st.text_input("GPA (out of 4)", key="app_info_gpa", placeholder="e.g. 3.8")
    st.text_input("SAT", key="app_info_sat", placeholder="e.g. 1450")
    st.text_input("ACT", key="app_info_act", placeholder="e.g. 32")
    st.text_input("Home State", key="app_info_home_state", placeholder="e.g. Virginia")
    st.text_input("Major Interest", key="app_info_major_interest", placeholder="e.g. Political Science")
    st.markdown("</div>", unsafe_allow_html=True)


def _truncate_text(text: str, limit: int = 2500) -> str:
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "\n... (truncated)"


def _mask_pii(text: str) -> str:
    masked_text = text
    for name, pattern in PII_PATTERNS.items():
        masked_text = pattern.sub(f"<{name}>", masked_text)
    masked_text = LOCATION_LINE_PATTERN.sub(
        lambda match: f"{match.group(1)}<location>", masked_text
    )
    return _mask_top_name(masked_text)


def _mask_top_name(text: str) -> str:
    lines = text.splitlines()
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) > 120:
            break
        if (
            re.fullmatch(r"[A-Za-z .,'\\-]+", stripped)
            and len(stripped.split()) <= 5
            and any(char.isalpha() for char in stripped)
        ):
            lines[index] = "<name>"
        break
    return "\n".join(lines)


def _hash_value(value: str) -> str:
    if not value:
        return ""
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _normalize_lookup_value(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold()


def _normalize_token(token: str) -> str:
    if token.endswith("ies") and len(token) > 4:
        return token[:-3] + "y"
    if token.endswith("s") and len(token) > 4:
        return token[:-1]
    return token


def _tokenize_for_match(value: str) -> set[str]:
    return {
        _normalize_token(token)
        for token in re.findall(r"[a-z0-9]+", value.casefold())
        if token not in MATCH_STOPWORDS
    }


def _clean_csv_value(value: str | None) -> str | None:
    if value is None:
        return None
    stripped = value.strip()
    return stripped or None


def _parse_optional_float(value: str | None) -> float | None:
    cleaned = _clean_csv_value(value)
    if cleaned is None or cleaned.casefold() in NULL_TOKENS:
        return None
    return float(cleaned)


def _parse_optional_int(value: str | None) -> int | None:
    parsed = _parse_optional_float(value)
    if parsed is None:
        return None
    return int(parsed)


def _parse_academic_year_start(value: str | None) -> int | None:
    cleaned = _clean_csv_value(value)
    if cleaned is None:
        return None

    match = re.match(r"^(\d{2})\s*-\s*(\d{2})$", cleaned)
    if not match:
        return None
    return 2000 + int(match.group(1))


def _normalize_supplemental_school_key(value: str | None) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", str(value or "").casefold())).strip()


def _get_dataset_signatures() -> dict[str, tuple[int, int]]:
    signatures: dict[str, tuple[int, int]] = {}
    for path in (DATASET_PATH, SUPPLEMENTAL_NET_PRICE_PATH):
        if not path.exists():
            continue
        stat = path.stat()
        signatures[str(path)] = (stat.st_size, stat.st_mtime_ns)
    return signatures


def _sqlite_exclusions_applied(connection: sqlite3.Connection) -> bool:
    if not EXCLUDED_SQL_UNITIDS:
        return True

    placeholders = ", ".join("?" for _ in EXCLUDED_SQL_UNITIDS)
    try:
        row = connection.execute(
            f"SELECT COUNT(*) FROM {SCORECARD_TABLE} WHERE UNITID IN ({placeholders})",
            tuple(sorted(EXCLUDED_SQL_UNITIDS)),
        ).fetchone()
    except sqlite3.Error:
        return False

    return bool(row) and int(row[0] or 0) == 0


def _metadata_matches_dataset(connection: sqlite3.Connection) -> bool:
    signatures = _get_dataset_signatures()
    primary_signature = signatures.get(str(DATASET_PATH))
    if primary_signature is None:
        return False

    try:
        rows = connection.execute(
            """
            SELECT dataset_path, dataset_size, dataset_mtime_ns
            FROM import_metadata
            """
        ).fetchall()
    except sqlite3.Error:
        return False

    metadata_rows = {
        str(row[0]): (int(row[1]), int(row[2]))
        for row in rows
    }
    for dataset_path, signature in signatures.items():
        if metadata_rows.get(dataset_path) != signature:
            return False

    return _sqlite_exclusions_applied(connection)


def _create_sqlite_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        f"""
        DROP TABLE IF EXISTS {SCORECARD_TABLE};
        DROP TABLE IF EXISTS {SUPPLEMENTAL_NET_PRICE_TABLE};
        DROP TABLE IF EXISTS import_metadata;

        CREATE TABLE {SCORECARD_TABLE} (
            UNITID INTEGER PRIMARY KEY,
            INSTNM TEXT NOT NULL,
            NORMALIZED_INSTNM TEXT NOT NULL,
            CITY TEXT,
            STABBR TEXT,
            CONTROL TEXT,
            LOCALE TEXT,
            CCBASIC TEXT,
            PREDDEG TEXT,
            HIGHDEG TEXT,
            ADM_RATE REAL,
            UGDS REAL,
            TUITIONFEE_IN REAL,
            TUITIONFEE_OUT REAL,
            PCTPELL REAL,
            PCTFLOAN REAL,
            DEBT_MDN REAL,
            OMAWDP8_ALL REAL,
            MD_EARN_WNE_P10 REAL
        );

        CREATE TABLE {SUPPLEMENTAL_NET_PRICE_TABLE} (
            row_id INTEGER PRIMARY KEY AUTOINCREMENT,
            school_name TEXT NOT NULL,
            normalized_name TEXT NOT NULL,
            year_label TEXT NOT NULL,
            year_start INTEGER NOT NULL,
            average_net_price REAL,
            net_price_income_0_30000 REAL,
            net_price_income_30001_48000 REAL,
            net_price_income_48001_75000 REAL,
            net_price_income_75001_110000 REAL,
            net_price_income_110001_plus REAL
        );

        CREATE TABLE import_metadata (
            dataset_path TEXT PRIMARY KEY,
            dataset_size INTEGER NOT NULL,
            dataset_mtime_ns INTEGER NOT NULL,
            imported_at TEXT NOT NULL
        );
        """
    )


def _build_sqlite_database() -> None:
    signatures = _get_dataset_signatures()
    primary_signature = signatures.get(str(DATASET_PATH))
    if primary_signature is None:
        return

    temp_path = SQLITE_PATH.with_suffix(".tmp")
    if temp_path.exists():
        temp_path.unlink()

    connection = sqlite3.connect(temp_path)
    try:
        connection.execute("PRAGMA journal_mode = OFF")
        connection.execute("PRAGMA synchronous = OFF")
        connection.execute("PRAGMA temp_store = MEMORY")
        connection.execute("PRAGMA cache_size = -50000")
        _create_sqlite_schema(connection)

        scorecard_insert_sql = f"""
            INSERT INTO {SCORECARD_TABLE} (
                UNITID,
                INSTNM,
                NORMALIZED_INSTNM,
                CITY,
                STABBR,
                CONTROL,
                LOCALE,
                CCBASIC,
                PREDDEG,
                HIGHDEG,
                ADM_RATE,
                UGDS,
                TUITIONFEE_IN,
                TUITIONFEE_OUT,
                PCTPELL,
                PCTFLOAN,
                DEBT_MDN,
                OMAWDP8_ALL,
                MD_EARN_WNE_P10
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """

        batch = []
        with DATASET_PATH.open(newline="", encoding="utf-8") as csv_file:
            reader = csv.DictReader(csv_file)
            for row in reader:
                unitid = _parse_optional_int(row["UNITID"])
                if unitid in EXCLUDED_SQL_UNITIDS:
                    continue

                batch.append(
                    (
                        unitid,
                        _clean_csv_value(row["INSTNM"]),
                        _normalize_supplemental_school_key(row["INSTNM"]),
                        _clean_csv_value(row["CITY"]),
                        _clean_csv_value(row["STABBR"]),
                        _clean_csv_value(row["CONTROL"]),
                        _clean_csv_value(row["LOCALE"]),
                        _clean_csv_value(row["CCBASIC"]),
                        _clean_csv_value(row["PREDDEG"]),
                        _clean_csv_value(row["HIGHDEG"]),
                        _parse_optional_float(row["ADM_RATE"]),
                        _parse_optional_float(row["UGDS"]),
                        _parse_optional_float(row["TUITIONFEE_IN"]),
                        _parse_optional_float(row["TUITIONFEE_OUT"]),
                        _parse_optional_float(row["PCTPELL"]),
                        _parse_optional_float(row["PCTFLOAN"]),
                        _parse_optional_float(row["DEBT_MDN"]),
                        _parse_optional_float(row["OMAWDP8_ALL"]),
                        _parse_optional_float(row["MD_EARN_WNE_P10"]),
                    )
                )
                if len(batch) >= SQLITE_IMPORT_BATCH_SIZE:
                    connection.executemany(scorecard_insert_sql, batch)
                    batch.clear()

            if batch:
                connection.executemany(scorecard_insert_sql, batch)

        if SUPPLEMENTAL_NET_PRICE_PATH.exists():
            supplemental_insert_sql = f"""
                INSERT INTO {SUPPLEMENTAL_NET_PRICE_TABLE} (
                    school_name,
                    normalized_name,
                    year_label,
                    year_start,
                    average_net_price,
                    net_price_income_0_30000,
                    net_price_income_30001_48000,
                    net_price_income_48001_75000,
                    net_price_income_75001_110000,
                    net_price_income_110001_plus
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """
            supplemental_batch = []
            with SUPPLEMENTAL_NET_PRICE_PATH.open(newline="", encoding="utf-8", errors="ignore") as csv_file:
                reader = csv.DictReader(csv_file)
                for row in reader:
                    school_name = _clean_csv_value(row.get("name"))
                    normalized_name = _normalize_supplemental_school_key(school_name)
                    year_label = _clean_csv_value(row.get("year"))
                    year_start = _parse_academic_year_start(year_label)
                    if not school_name or not normalized_name or not year_label or year_start is None:
                        continue

                    supplemental_batch.append(
                        (
                            school_name,
                            normalized_name,
                            year_label,
                            year_start,
                            _parse_optional_float(row.get("averageNetPrice")),
                            _parse_optional_float(row.get("netPriceIncome0to30000")),
                            _parse_optional_float(row.get("netPriceIncome30001to48000")),
                            _parse_optional_float(row.get("netPriceIncome48001to75000")),
                            _parse_optional_float(row.get("netPriceIncome75001to110000")),
                            _parse_optional_float(row.get("netPriceIncome110001")),
                        )
                    )
                    if len(supplemental_batch) >= SQLITE_IMPORT_BATCH_SIZE:
                        connection.executemany(supplemental_insert_sql, supplemental_batch)
                        supplemental_batch.clear()

                if supplemental_batch:
                    connection.executemany(supplemental_insert_sql, supplemental_batch)

        connection.executescript(
            f"""
            CREATE INDEX idx_scorecard_instnm
            ON {SCORECARD_TABLE} (INSTNM);

            CREATE INDEX idx_scorecard_normalized_instnm
            ON {SCORECARD_TABLE} (NORMALIZED_INSTNM);

            CREATE INDEX idx_scorecard_state
            ON {SCORECARD_TABLE} (STABBR);

            CREATE INDEX idx_scorecard_control
            ON {SCORECARD_TABLE} (CONTROL);

            CREATE INDEX idx_scorecard_earnings
            ON {SCORECARD_TABLE} (MD_EARN_WNE_P10);

            CREATE INDEX idx_supplemental_net_price_name
            ON {SUPPLEMENTAL_NET_PRICE_TABLE} (normalized_name);

            CREATE INDEX idx_supplemental_net_price_year
            ON {SUPPLEMENTAL_NET_PRICE_TABLE} (year_start);
            """
        )

        connection.executemany(
            """
            INSERT INTO import_metadata (
                dataset_path,
                dataset_size,
                dataset_mtime_ns,
                imported_at
            )
            VALUES (?, ?, ?, CURRENT_TIMESTAMP)
            """,
            [
                (dataset_path, signature[0], signature[1])
                for dataset_path, signature in sorted(signatures.items())
            ],
        )
        connection.commit()
    finally:
        connection.close()

    os.replace(temp_path, SQLITE_PATH)


def _ensure_sqlite_database() -> None:
    if SQLITE_PATH.exists() and not DATASET_PATH.exists():
        return

    if not DATASET_PATH.exists():
        return

    if SQLITE_PATH.exists():
        try:
            with sqlite3.connect(SQLITE_PATH) as connection:
                if _metadata_matches_dataset(connection):
                    return
        except sqlite3.Error:
            logger.warning("Existing SQLite database could not be read. Rebuilding.")

    _build_sqlite_database()


def _log_resume_upload(filename: str, preview: str) -> None:
    logger.info(
        "Resume upload recorded filename_hash=%s preview=%s",
        _hash_value(filename),
        _mask_pii(preview)[:200],
    )


def _extract_pdf_text(buffer: BytesIO) -> str:
    buffer.seek(0)
    text_chunks = []
    with pdfplumber.open(buffer) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def _extract_docx_text(buffer: BytesIO) -> str:
    buffer.seek(0)
    document = Document(buffer)
    return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()


def _extract_plain_text(buffer: BytesIO) -> str:
    buffer.seek(0)
    return buffer.read().decode("utf-8", errors="ignore").strip()


def _parse_resume(uploaded_file) -> str:
    try:
        uploaded_file.seek(0)
    except Exception:
        pass
    content = uploaded_file.read()
    buffer = BytesIO(content)
    filename = uploaded_file.name.lower()
    if filename.endswith(".pdf"):
        return _extract_pdf_text(buffer)
    if filename.endswith(".docx"):
        return _extract_docx_text(buffer)
    return _extract_plain_text(buffer)


def _process_uploaded_resume(uploaded_resume) -> None:
    try:
        parsed_text = _parse_resume(uploaded_resume)
        if not parsed_text:
            st.warning("Could not extract meaningful text from that file.")
            return

        redacted_text = _mask_pii(parsed_text)
        st.session_state.resume_text = redacted_text
        st.session_state.resume_filename = uploaded_resume.name
        _log_resume_upload(uploaded_resume.name, _truncate_text(redacted_text, 500))
        st.success(f"Loaded resume: {uploaded_resume.name}")
    except Exception as err:
        st.error(f"Resume parsing failed: {err}")


def _clear_resume_state() -> None:
    st.session_state.resume_text = ""
    st.session_state.resume_filename = ""


def _chat_with_model(deployment_id: str, messages: list[dict]) -> str:
    response = openai.chat.completions.create(
        model=deployment_id,
        messages=messages,
        max_completion_tokens=AZURE_MAX_RESPONSE_TOKENS,
        reasoning_effort=AZURE_REASONING_EFFORT,
    )
    choice = response.choices[0] if response.choices else None
    message = getattr(choice, "message", None)

    reply = _coerce_model_content_to_text(getattr(message, "content", None))
    if reply:
        return reply.strip()

    refusal = _coerce_model_content_to_text(getattr(message, "refusal", None))
    if refusal:
        return f"The Azure model declined to answer directly: {refusal}"

    diagnostic = _describe_empty_chat_response(response)
    logger.warning(diagnostic)
    return diagnostic


def _coerce_model_content_to_text(content) -> str:
    if content is None:
        return ""
    if isinstance(content, str):
        return content
    if hasattr(content, "model_dump"):
        return _coerce_model_content_to_text(content.model_dump())
    if isinstance(content, list):
        parts = []
        for item in content:
            text_value = _coerce_model_content_to_text(item)
            if text_value:
                parts.append(text_value)
        return "".join(parts).strip()
    if isinstance(content, dict):
        if content.get("type") in {"text", "output_text"}:
            text_value = content.get("text")
            if text_value is not None:
                return _coerce_model_content_to_text(text_value)
        for key in ("text", "value", "content"):
            if key in content and content[key] is not None:
                text_value = _coerce_model_content_to_text(content[key])
                if text_value:
                    return text_value
        return ""
    for attr_name in ("text", "value", "content"):
        attr_value = getattr(content, attr_name, None)
        if attr_value is not None:
            text_value = _coerce_model_content_to_text(attr_value)
            if text_value:
                return text_value
    return str(content).strip()


def _describe_empty_chat_response(response) -> str:
    try:
        choice = response.choices[0] if response.choices else None
        if choice is None:
            return "Azure returned no choices for this prompt."

        finish_reason = getattr(choice, "finish_reason", None) or "unknown"
        message = getattr(choice, "message", None)
        refusal = _coerce_model_content_to_text(getattr(message, "refusal", None))

        if refusal:
            return f"The Azure model declined to answer directly: {refusal}"
        if finish_reason == "content_filter":
            return "Azure blocked the response because it triggered content filtering."
        if finish_reason == "length":
            return "Azure stopped before returning visible text because the reply hit the completion limit."

        if hasattr(response, "model_dump_json"):
            logger.warning(
                "Empty Azure chat completion payload: %s",
                _truncate_text(response.model_dump_json(indent=2), 3000),
            )

        return f"Azure returned no visible assistant text for this prompt (finish_reason={finish_reason})."
    except Exception:
        logger.exception("Failed to inspect empty Azure response.")
        return "Azure returned no visible assistant text for this prompt."


def _word_count(text: str) -> int:
    return len(re.findall(r"\b\S+\b", text))


def _trim_text_to_word_limit(text: str, max_words: int) -> str:
    normalized = re.sub(r"\s+", " ", str(text or "").strip())
    if not normalized:
        return ""

    words = list(re.finditer(r"\S+", normalized))
    if len(words) <= max_words:
        return normalized

    cutoff_index = words[max_words - 1].end()
    candidate = normalized[:cutoff_index].strip()

    min_boundary_word_index = max(0, min(len(words) - 1, int(max_words * 0.6) - 1))
    min_boundary_index = words[min_boundary_word_index].end() if words else 0
    sentence_matches = [
        match
        for match in re.finditer(r"[.!?](?:[\"')\]]+)?(?=\s|$)", candidate)
        if match.end() >= min_boundary_index
    ]
    if sentence_matches:
        return candidate[: sentence_matches[-1].end()].strip()

    if candidate.endswith((".", "!", "?")):
        return candidate
    return candidate.rstrip(",;:-") + "..."


def _compact_display_text(
    text: str,
    *,
    max_words: int,
    max_sentences: int,
    max_paragraphs: int,
    overflow_words: int = 0,
) -> str:
    normalized = re.sub(r"\n{3,}", "\n\n", str(text or "").strip())
    if not normalized:
        return ""

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    selected_paragraphs: list[str] = []
    words_used = 0
    sentences_used = 0
    hard_word_limit = max_words + max(0, overflow_words)
    overflow_sentence_used = False

    for paragraph in paragraphs:
        paragraph_sentences = _split_into_sentences(paragraph)
        if not paragraph_sentences:
            paragraph_sentences = [paragraph]

        chosen_sentences: list[str] = []
        for sentence in paragraph_sentences:
            sentence_words = _word_count(sentence)
            projected_words = words_used + sentence_words

            if not chosen_sentences and not selected_paragraphs and sentence_words > hard_word_limit:
                trimmed_sentence = _trim_text_to_word_limit(sentence, hard_word_limit)
                chosen_sentences.append(trimmed_sentence)
                words_used = _word_count(trimmed_sentence)
                sentences_used += 1
                overflow_sentence_used = words_used > max_words
                break

            if sentences_used >= max_sentences:
                break

            if projected_words > max_words:
                can_finish_with_one_more_sentence = (
                    overflow_words > 0
                    and not overflow_sentence_used
                    and projected_words <= hard_word_limit
                )
                if can_finish_with_one_more_sentence:
                    chosen_sentences.append(sentence)
                    words_used = projected_words
                    sentences_used += 1
                    overflow_sentence_used = True
                break

            chosen_sentences.append(sentence)
            words_used = projected_words
            sentences_used += 1

        if chosen_sentences:
            selected_paragraphs.append(" ".join(chosen_sentences))
        if (
            len(selected_paragraphs) >= max_paragraphs
            or sentences_used >= max_sentences
            or words_used >= hard_word_limit
        ):
            break

    if not selected_paragraphs:
        return _trim_text_to_word_limit(normalized, hard_word_limit)

    compacted = "\n\n".join(selected_paragraphs[:max_paragraphs]).strip()
    if _word_count(compacted) > hard_word_limit:
        return _trim_text_to_word_limit(compacted, hard_word_limit)
    return compacted
def _stream_text_chunks(text: str, chunk_size: int = 14):
    if not text:
        return
    words = text.split(" ")
    chunk: list[str] = []
    for word in words:
        chunk.append(word)
        if len(chunk) >= chunk_size:
            yield " ".join(chunk) + " "
            chunk = []
    if chunk:
        yield " ".join(chunk)


def _stream_chat_with_model(deployment_id: str, messages: list[dict]):
    try:
        response = openai.chat.completions.create(
            model=deployment_id,
            messages=messages,
            stream=True,
            max_completion_tokens=AZURE_MAX_RESPONSE_TOKENS,
            reasoning_effort=AZURE_REASONING_EFFORT,
        )
        streamed_parts: list[str] = []
        for chunk in response:
            if not chunk.choices:
                continue
            delta = getattr(chunk.choices[0].delta, "content", None)
            if delta:
                chunk_text = _coerce_model_content_to_text(delta)
                if chunk_text:
                    streamed_parts.append(chunk_text)
        if streamed_parts:
            reply = "".join(streamed_parts).strip()
            if reply:
                yield from _stream_text_chunks(reply)
            return
    except Exception:
        logger.exception("Streaming chat fell back to a non-streaming response.")

    reply = _chat_with_model(deployment_id, messages)
    if reply.strip():
        yield from _stream_text_chunks(reply)
        return

    yield "Azure returned an empty response for this prompt."


def _metric_expression(metric_key: str) -> str:
    return METRIC_DEFINITIONS[metric_key]["expression"]


def _projected_metric_expression(metric_key: str) -> str:
    if metric_key == "net_price":
        return "net_price"
    return _metric_expression(metric_key)


def _metric_format(metric_key: str) -> str:
    return METRIC_DEFINITIONS[metric_key]["format"]


def _metric_axis_format(metric_key: str) -> str:
    metric_format = _metric_format(metric_key)
    if metric_format == "currency":
        return "$,.0f"
    if metric_format == "percent":
        return ".0%"
    return ",.0f"


def _format_currency(value: float | None) -> str:
    if value is None:
        return "n/a"
    return f"${value:,.0f}"


def _format_count(value: float | None) -> str:
    if value is None:
        return "n/a"
    return f"{int(value):,}"


def _format_percent(value: float | None) -> str:
    if value is None:
        return "n/a"
    return f"{value * 100:.1f}%"


def _format_metric_value(metric_key: str, value: float | None) -> str:
    metric_format = _metric_format(metric_key)
    if metric_format == "currency":
        return _format_currency(value)
    if metric_format == "percent":
        return _format_percent(value)
    return _format_count(value)


def _format_location(city: str | None, state: str | None) -> str:
    parts = [part for part in [city, state] if part]
    return ", ".join(parts) if parts else "n/a"


def _coerce_float(value) -> float | None:
    if value is None:
        return None
    return float(value)


def _row_to_record(row: sqlite3.Row, metric_key: str) -> dict:
    control_code = str(row["CONTROL"]).strip() if row["CONTROL"] is not None else ""
    return {
        "unitid": row["UNITID"],
        "institution_name": row["INSTNM"],
        "city": row["CITY"],
        "state": row["STABBR"],
        "location": _format_location(row["CITY"], row["STABBR"]),
        "control_code": control_code,
        "control": CONTROL_LABELS.get(control_code, "Unknown"),
        "locale_code": row["LOCALE"],
        "carnegie_code": row["CCBASIC"],
        "preddeg_code": row["PREDDEG"],
        "highdeg_code": row["HIGHDEG"],
        "metric_key": metric_key,
        "metric_value": _coerce_float(row["metric_value"]),
        "admission_rate": _coerce_float(row["ADM_RATE"]),
        "undergrad_size": _coerce_float(row["UGDS"]),
        "net_price": _coerce_float(row["net_price"]),
        "net_price_source": "supplemental_history" if row["net_price"] is not None else "",
        "tuition_in": _coerce_float(row["TUITIONFEE_IN"]),
        "tuition_out": _coerce_float(row["TUITIONFEE_OUT"]),
        "pell_share": _coerce_float(row["PCTPELL"]),
        "loan_share": _coerce_float(row["PCTFLOAN"]),
        "debt": _coerce_float(row["DEBT_MDN"]),
        "completion_rate": _coerce_float(row["OMAWDP8_ALL"]),
        "earnings": _coerce_float(row["MD_EARN_WNE_P10"]),
    }


def _build_supplemental_history_lookup_keys(school_name: str) -> tuple[str, ...]:
    normalized_name = _normalize_supplemental_school_key(school_name)
    keys: list[str] = []
    seen: set[str] = set()

    def add_key(value: str | None) -> None:
        normalized = _normalize_supplemental_school_key(value)
        if not normalized or normalized in seen:
            return
        seen.add(normalized)
        keys.append(normalized)

    add_key(school_name)
    for alias in sorted(_generate_school_aliases(school_name), key=len, reverse=True):
        add_key(alias)

    return tuple(keys)


def _fetch_supplemental_net_price_history(
    connection: sqlite3.Connection,
    school_name: str,
    limit_years: int = SUPPLEMENTAL_NET_PRICE_CHART_YEARS,
) -> tuple[list[dict], str | None]:
    lookup_keys = _build_supplemental_history_lookup_keys(school_name)
    if not lookup_keys:
        return [], None

    sql = f"""
        SELECT
            MIN(school_name) AS school_name,
            year_label,
            year_start,
            AVG(average_net_price) AS average_net_price,
            AVG(net_price_income_0_30000) AS net_price_income_0_30000,
            AVG(net_price_income_30001_48000) AS net_price_income_30001_48000,
            AVG(net_price_income_48001_75000) AS net_price_income_48001_75000,
            AVG(net_price_income_75001_110000) AS net_price_income_75001_110000,
            AVG(net_price_income_110001_plus) AS net_price_income_110001_plus
        FROM {SUPPLEMENTAL_NET_PRICE_TABLE}
        WHERE normalized_name = ?
        GROUP BY year_label, year_start
        ORDER BY year_start DESC
        LIMIT ?
    """
    for lookup_key in lookup_keys:
        rows = connection.execute(sql, (lookup_key, limit_years)).fetchall()
        if not rows:
            continue
        history_rows = [
            {
                "school_name": row["school_name"],
                "year_label": row["year_label"],
                "year_start": int(row["year_start"]),
                "average_net_price": _coerce_float(row["average_net_price"]),
                "net_price_income_0_30000": _coerce_float(row["net_price_income_0_30000"]),
                "net_price_income_30001_48000": _coerce_float(row["net_price_income_30001_48000"]),
                "net_price_income_48001_75000": _coerce_float(row["net_price_income_48001_75000"]),
                "net_price_income_75001_110000": _coerce_float(row["net_price_income_75001_110000"]),
                "net_price_income_110001_plus": _coerce_float(row["net_price_income_110001_plus"]),
            }
            for row in rows
        ]
        history_rows.sort(key=lambda item: item["year_start"])
        return history_rows, lookup_key

    return [], None


def _apply_supplemental_net_price_fallback(
    connection: sqlite3.Connection,
    rows: list[dict],
) -> list[dict]:
    if not rows:
        return rows

    history_cache: dict[str, list[dict]] = {}
    for row in rows:
        if row.get("net_price") is not None:
            row["net_price_source"] = "supplemental_history"
            continue

        school_name = str(row.get("institution_name", "")).strip()
        if not school_name:
            continue

        history_rows = history_cache.get(school_name)
        if history_rows is None:
            history_rows, _ = _fetch_supplemental_net_price_history(
                connection,
                school_name,
                limit_years=SUPPLEMENTAL_NET_PRICE_CHART_YEARS,
            )
            history_cache[school_name] = history_rows

        latest_average = history_rows[-1]["average_net_price"] if history_rows else None
        if latest_average is None:
            continue

        row["net_price"] = latest_average
        row["net_price_source"] = "supplemental_history"
        if row.get("metric_key") == "net_price" and row.get("metric_value") is None:
            row["metric_value"] = latest_average

    return rows


def _load_school_names() -> list[str]:
    if not SQLITE_PATH.exists() and not DATASET_PATH.exists():
        return []

    _ensure_sqlite_database()
    with sqlite3.connect(SQLITE_PATH) as connection:
        names = [
            row[0]
            for row in connection.execute(
                f"""
                SELECT INSTNM
                FROM {SCORECARD_TABLE}
                WHERE INSTNM IS NOT NULL
                ORDER BY LENGTH(INSTNM) DESC, INSTNM
                """
            ).fetchall()
        ]
    return names


def _generate_school_aliases(school_name: str) -> set[str]:
    normalized = _normalize_lookup_value(school_name)
    if not normalized:
        return set()

    aliases = {normalized}
    tokens = normalized.split()

    if normalized.startswith("the "):
        aliases.add(normalized[4:])

    if len(tokens) > 1 and tokens[-1] in INSTITUTION_TYPE_TOKENS:
        aliases.add(" ".join(tokens[:-1]))

    if normalized.startswith("university of "):
        aliases.add(normalized.removeprefix("university of ").strip())
    if normalized.startswith("the university of "):
        aliases.add(normalized.removeprefix("the university of ").strip())

    meaningful_tokens = [
        token
        for token in tokens
        if token not in SCHOOL_NAME_FILLER_TOKENS and token not in INSTITUTION_TYPE_TOKENS
    ]
    if len(meaningful_tokens) >= 2:
        aliases.add(" ".join(meaningful_tokens))
    elif len(meaningful_tokens) == 1 and len(meaningful_tokens[0]) >= 6:
        aliases.add(meaningful_tokens[0])

    blocked_single_word_aliases = BLOCKED_GEOGRAPHIC_ALIASES | {
        abbreviation.casefold() for abbreviation in STATE_ABBRS
    }

    cleaned_aliases = set()
    for alias in aliases:
        candidate = alias.strip()
        if len(candidate) < 3:
            continue
        if candidate in BLOCKED_GEOGRAPHIC_ALIASES:
            continue
        if " " not in candidate and candidate in blocked_single_word_aliases:
            continue
        cleaned_aliases.add(candidate)
    return cleaned_aliases


def _generate_school_acronyms(school_name: str) -> set[str]:
    normalized = _normalize_lookup_value(school_name)
    tokens = normalized.split()
    acronym_tokens = [
        token
        for token in tokens
        if token not in SCHOOL_NAME_FILLER_TOKENS
    ]
    acronym = "".join(token[0] for token in acronym_tokens if token)
    if 3 <= len(acronym) <= 6:
        return {acronym}
    return set()


def _school_name_priority(school_name: str) -> tuple[int, int]:
    normalized = _normalize_lookup_value(school_name)
    score = 0
    if " university " in f" {normalized} " or normalized.startswith("university ") or normalized.endswith(" university"):
        score += 4
    if " college " in f" {normalized} " or normalized.startswith("college ") or normalized.endswith(" college"):
        score += 2
    if " institute " in f" {normalized} " or normalized.startswith("institute ") or normalized.endswith(" institute"):
        score += 1
    return score, len(normalized)


def _load_school_alias_lookup() -> dict[str, str]:
    alias_to_schools: dict[str, set[str]] = {}
    for school in _load_school_names():
        for alias in _generate_school_aliases(school):
            alias_to_schools.setdefault(alias, set()).add(school)

    resolved: dict[str, str] = {}
    for alias, schools in alias_to_schools.items():
        ranked = sorted(
            ((_school_name_priority(school), school) for school in schools),
            key=lambda item: (item[0][0], item[0][1], item[1]),
            reverse=True,
        )
        if not ranked:
            continue
        if len(ranked) == 1 or ranked[0][0] > ranked[1][0]:
            resolved[alias] = ranked[0][1]
    return resolved


def _load_school_acronym_lookup() -> dict[str, str]:
    acronym_to_schools: dict[str, set[str]] = {}
    for school in _load_school_names():
        for acronym in _generate_school_acronyms(school):
            acronym_to_schools.setdefault(acronym, set()).add(school)

    resolved: dict[str, str] = {}
    for acronym, schools in acronym_to_schools.items():
        ranked = sorted(
            ((_school_name_priority(school), school) for school in schools),
            key=lambda item: (item[0][0], item[0][1], item[1]),
            reverse=True,
        )
        if not ranked:
            continue
        if len(ranked) == 1 or ranked[0][0] > ranked[1][0]:
            resolved[acronym] = ranked[0][1]
    return resolved


def _extract_requested_limit(question: str, default: int = DEFAULT_RESULT_LIMIT) -> int:
    match = re.search(r"\b(?:top|bottom|first)\s+(\d{1,2})\b", question, re.IGNORECASE)
    if match:
        return max(1, min(int(match.group(1)), MAX_RESULT_LIMIT))
    return default


def _match_schools(question: str, schools: list[str], maximum: int = 8) -> list[str]:
    question_normalized = _normalize_lookup_value(question)
    alias_lookup = _load_school_alias_lookup()
    acronym_lookup = _load_school_acronym_lookup()
    best_matches: dict[str, tuple[int, int]] = {}

    for alias, school in alias_lookup.items():
        match = re.search(rf"(?<!\w){re.escape(alias)}(?!\w)", question_normalized)
        if not match:
            continue
        score = (match.start(), -len(alias))
        if school not in best_matches or score < best_matches[school]:
            best_matches[school] = score

    uppercase_tokens = set(re.findall(r"\b[A-Z]{3,6}\b", question))
    for token in uppercase_tokens:
        school = acronym_lookup.get(token.casefold())
        if not school:
            continue
        match = re.search(rf"\b{re.escape(token)}\b", question)
        if not match:
            continue
        score = (match.start(), -len(token))
        if school not in best_matches or score < best_matches[school]:
            best_matches[school] = score

    ordered_matches = sorted(
        ((position, alias_length, school) for school, (position, alias_length) in best_matches.items()),
        key=lambda item: (item[0], item[1], item[2]),
    )
    return [school for _, _, school in ordered_matches[:maximum]]


def _metric_keyword_checks() -> list[tuple[list[str], str]]:
    return [
        (
            [
                "net price",
                "net cost",
                "cost",
                "costs",
                "cost after aid",
                "after aid",
                "after grants",
                "affordability",
                "affordable",
                "cheapest",
            ],
            "net_price",
        ),
        (["out-of-state tuition", "out of state tuition"], "tuition_out"),
        (["in-state tuition", "in state tuition"], "tuition_in"),
        (["tuition"], "tuition_in"),
        (["median debt", "student debt", "loan debt", "debt"], "debt"),
        (
            [
                "acceptance rate",
                "admission rate",
                "admissions rate",
                "admit rate",
                "get into",
                "get in",
                "selective",
                "hard to get into",
                "easy to get into",
            ],
            "admission_rate",
        ),
        (["completion", "graduate", "graduation", "award rate"], "completion_rate"),
        (["pell"], "pell_share"),
        (["loan share", "students taking loans", "borrowers", "federal loans"], "loan_share"),
        (["how big", "how large", "enrollment", "undergrads", "undergraduate", "student body", "size", "big", "large"], "undergrad_size"),
        (["earnings", "salary", "salaries", "income"], "earnings"),
    ]


def _extract_metric_mentions(question: str) -> list[str]:
    lower_question = question.casefold()
    matches: list[tuple[int, int, str]] = []
    for keywords, metric_key in _metric_keyword_checks():
        best_position: tuple[int, int] | None = None
        for keyword in keywords:
            position = lower_question.find(keyword)
            if position < 0:
                continue
            score = (position, -len(keyword))
            if best_position is None or score < best_position:
                best_position = score
        if best_position is not None:
            matches.append((best_position[0], best_position[1], metric_key))

    matches.sort()
    ordered_metrics: list[str] = []
    for _, _, metric_key in matches:
        if metric_key not in ordered_metrics:
            ordered_metrics.append(metric_key)
    return ordered_metrics


def _resolve_metric(question: str) -> tuple[str, bool, list[str]]:
    lower_question = question.casefold()
    assumptions: list[str] = []

    best_match: tuple[int, int, str, str] | None = None
    for keywords, metric_key in _metric_keyword_checks():
        for keyword in keywords:
            position = lower_question.find(keyword)
            if position < 0:
                continue
            score = (position, -len(keyword), metric_key, keyword)
            if best_match is None or score < best_match:
                best_match = score

    if best_match:
        _, _, metric_key, keyword = best_match
        if metric_key == "tuition_in" and keyword == "tuition" and "out-of-state" not in lower_question and "in-state" not in lower_question:
            assumptions.append(
                "The question said `tuition` without specifying in-state or out-of-state, so the results default to in-state tuition."
            )
        return metric_key, True, assumptions

    assumptions.append(
        "No specific institution metric was requested, so the school profile defaults to median earnings and also shows other institutional stats in the table."
    )
    return DEFAULT_METRIC_KEY, False, assumptions


def _extract_benchmark_stat(question: str) -> str:
    lower_question = question.casefold()
    candidates = []
    for token, stat in (("average", "average"), ("mean", "average"), ("median", "median")):
        position = lower_question.rfind(token)
        if position >= 0:
            candidates.append((position, stat))

    if not candidates:
        return "average"

    candidates.sort()
    return candidates[-1][1]


def _looks_like_unmatched_school_request(
    question: str,
    matched_schools: list[str],
    wants_ranked_results: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
) -> bool:
    if matched_schools or state_filter or control_filter:
        return False

    lower_question = question.casefold()
    if wants_ranked_results and not re.search(r"\b(university|college|institute|school)\b", lower_question):
        return False

    if re.search(
        r"\b[A-Z][A-Za-z&.'-]+(?:\s+[A-Z][A-Za-z&.'-]+){0,5}\s+(?:University|College|Institute|School)\b",
        question,
    ):
        return True

    uppercase_tokens = set(re.findall(r"\b[A-Z]{3,6}\b", question))
    if uppercase_tokens:
        return True

    return False


def _extract_state_filter(question: str, ignore_phrases: list[str] | None = None) -> str | None:
    lower_question = question.casefold()
    for phrase in ignore_phrases or []:
        if phrase:
            lower_question = lower_question.replace(phrase.casefold(), " ")
    for state_name, abbreviation in STATE_NAME_TO_ABBR.items():
        if re.search(rf"\b{re.escape(state_name)}\b", lower_question):
            return abbreviation

    uppercase_tokens = set(re.findall(r"\b[A-Z]{2}\b", question))
    for token in uppercase_tokens:
        if token in STATE_ABBRS:
            return token
    return None


def _normalize_control_codes(control_codes: tuple[str, ...] | list[str] | None) -> tuple[str, ...] | None:
    if not control_codes:
        return None
    normalized = tuple(sorted({str(code).strip() for code in control_codes if str(code).strip()}))
    return normalized or None


def _format_control_scope(control_codes: tuple[str, ...] | list[str] | None) -> str | None:
    normalized = _normalize_control_codes(control_codes)
    if not normalized:
        return None
    if normalized in CONTROL_SCOPE_LABELS:
        return CONTROL_SCOPE_LABELS[normalized]
    return ", ".join(CONTROL_LABELS.get(code, code) for code in normalized)


def _apply_control_filter(
    where_clauses: list[str],
    params: list[object],
    control_codes: tuple[str, ...] | list[str] | None,
) -> None:
    normalized = _normalize_control_codes(control_codes)
    if not normalized:
        return
    if len(normalized) == 1:
        where_clauses.append("CONTROL = ?")
        params.append(normalized[0])
        return
    placeholders = ", ".join("?" for _ in normalized)
    where_clauses.append(f"CONTROL IN ({placeholders})")
    params.extend(normalized)


def _extract_control_filter(question: str) -> tuple[str, ...] | None:
    lower_question = question.casefold()
    if "private nonprofit" in lower_question or "private non-profit" in lower_question or "nonprofit" in lower_question:
        return ("2",)
    if "for-profit" in lower_question or "for profit" in lower_question:
        return ("3",)
    if "private" in lower_question:
        return ("2", "3")
    if "public" in lower_question:
        return ("1",)
    return None


def _wants_distribution_graph(question: str) -> bool:
    lower_question = question.casefold()
    direct_phrases = [
        "where does",
        "where do",
        "fall compared",
        "falls compared",
        "fall relative",
        "falls relative",
        "stack up",
        "stack against",
        "compared to similar schools",
    ]
    if any(phrase in lower_question for phrase in direct_phrases):
        return True

    benchmark_phrases = ["compared to", "compare to", "compared with", "against", "relative to", "versus"]
    cohort_tokens = [
        "every",
        "all",
        "average",
        "mean",
        "median",
        "nationwide",
        "nationally",
        "country",
        "america",
        "united states",
        "peer",
        "similar schools",
    ]
    return any(phrase in lower_question for phrase in benchmark_phrases) and any(
        token in lower_question for token in cohort_tokens
    )


def _mentions_national_scope(question: str) -> bool:
    lower_question = question.casefold()
    return any(
        phrase in lower_question
        for phrase in [
            "country",
            "america",
            "nationwide",
            "nationally",
            "united states",
            "across the country",
            "in the nation",
            "in america",
        ]
    )


def _extract_sort_direction(question: str, metric_key: str) -> str:
    lower_question = question.casefold()

    if metric_key == "admission_rate":
        if "hardest to get into" in lower_question or "most selective" in lower_question:
            return "ASC"
        if "easiest to get into" in lower_question:
            return "DESC"

    if any(
        token in lower_question
        for token in ["lowest", "least", "smallest", "fewest", "cheapest", "bottom"]
    ):
        return "ASC"

    if any(token in lower_question for token in ["highest", "largest", "most"]):
        return "DESC"

    return METRIC_DEFINITIONS[metric_key]["default_direction"]


def _wants_ranked_results(question: str) -> bool:
    lower_question = question.casefold()
    return any(
        token in lower_question
        for token in [
            "top",
            "bottom",
            "highest",
            "lowest",
            "largest",
            "smallest",
            "best",
            "cheapest",
            "rank",
            "show me",
        ]
    )


def _sql_literal(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def _render_sql_with_params(query: str, params: list[object] | tuple[object, ...] | None = None) -> str:
    cleaned_query = query.strip()
    if not params:
        return cleaned_query

    parts = cleaned_query.split("?")
    if len(parts) - 1 != len(params):
        params_sql = ", ".join(_sql_literal(param) for param in params)
        return f"{cleaned_query}\n\n-- params: {params_sql}"

    rendered = parts[0]
    for index, param in enumerate(params):
        rendered += _sql_literal(param) + parts[index + 1]
    return rendered


def _build_fetch_institutions_sql(
    schools: list[str],
    metric_key: str,
) -> tuple[str, list[object]]:
    placeholders = ", ".join("?" for _ in schools)
    metric_expression = _projected_metric_expression(metric_key)
    sql = f"""
        SELECT base.*, {metric_expression} AS metric_value
        FROM ({COMMON_SELECT_SQL}) AS base
        WHERE INSTNM IN ({placeholders})
    """
    return sql, list(schools)


def _build_ranked_institutions_sql(
    metric_key: str,
    limit: int,
    sort_direction: str,
    state_filter: str | None = None,
    control_filter: tuple[str, ...] | None = None,
) -> tuple[str, list[object]]:
    metric_expression = _projected_metric_expression(metric_key)
    where_clauses = [f"{metric_expression} IS NOT NULL"]
    params: list[object] = []

    if state_filter:
        where_clauses.append("STABBR = ?")
        params.append(state_filter)
    _apply_control_filter(where_clauses, params, control_filter)
    params.append(limit)

    sql = f"""
        SELECT base.*, {metric_expression} AS metric_value
        FROM ({COMMON_SELECT_SQL}) AS base
        WHERE {' AND '.join(where_clauses)}
        ORDER BY metric_value {sort_direction}, INSTNM
        LIMIT ?
    """
    return sql, params


def _build_distribution_sql(
    metric_key: str,
    state_filter: str | None = None,
    states: tuple[str, ...] | None = None,
    control_codes: tuple[str, ...] | None = None,
) -> tuple[str, list[object]]:
    metric_expression = _metric_expression(metric_key)
    where_clauses = [f"{metric_expression} IS NOT NULL"]
    params: list[object] = []

    if state_filter:
        where_clauses.append("STABBR = ?")
        params.append(state_filter)
    if states:
        placeholders = ", ".join("?" for _ in states)
        where_clauses.append(f"STABBR IN ({placeholders})")
        params.extend(states)
    _apply_control_filter(where_clauses, params, control_codes)

    sql = f"""
        SELECT INSTNM, {metric_expression} AS metric_value
        FROM {SCORECARD_TABLE}
        WHERE {' AND '.join(where_clauses)}
        ORDER BY metric_value, INSTNM
    """
    return sql, params


def _build_benchmark_sql(
    metric_key: str,
    benchmark_stat: str,
    states: tuple[str, ...] | None = None,
    control_codes: tuple[str, ...] | None = None,
) -> tuple[str, list[object]]:
    metric_expression = _metric_expression(metric_key)
    where_clauses = [f"{metric_expression} IS NOT NULL"]
    params: list[object] = []

    if states:
        placeholders = ", ".join("?" for _ in states)
        where_clauses.append(f"STABBR IN ({placeholders})")
        params.extend(states)
    _apply_control_filter(where_clauses, params, control_codes)

    if benchmark_stat == "median":
        sql = f"""
            WITH benchmark_values AS (
                SELECT {metric_expression} AS metric_value
                FROM {SCORECARD_TABLE}
                WHERE {' AND '.join(where_clauses)}
            ),
            ordered_values AS (
                SELECT
                    metric_value,
                    ROW_NUMBER() OVER (ORDER BY metric_value) AS row_num,
                    COUNT(*) OVER () AS total_count
                FROM benchmark_values
            )
            SELECT AVG(metric_value)
            FROM ordered_values
            WHERE row_num IN ((total_count + 1) / 2, (total_count + 2) / 2)
        """
    else:
        sql = f"""
            SELECT AVG(metric_value)
            FROM (
                SELECT {metric_expression} AS metric_value
                FROM {SCORECARD_TABLE}
                WHERE {' AND '.join(where_clauses)}
            )
        """
    return sql, params


def _fetch_institutions_by_name(
    connection: sqlite3.Connection,
    schools: list[str],
    metric_key: str,
) -> list[dict]:
    if not schools:
        return []

    sql, params = _build_fetch_institutions_sql(schools, metric_key)
    rows = connection.execute(sql, params).fetchall()

    by_name = {row["INSTNM"]: _row_to_record(row, metric_key) for row in rows}
    ordered = []
    for school in schools:
        record = by_name.get(school)
        if record is not None:
            ordered.append(record)
    return _apply_supplemental_net_price_fallback(connection, ordered)


def _query_ranked_institutions(
    connection: sqlite3.Connection,
    metric_key: str,
    limit: int,
    sort_direction: str,
    state_filter: str | None = None,
    control_filter: tuple[str, ...] | None = None,
) -> list[dict]:
    sql, params = _build_ranked_institutions_sql(
        metric_key=metric_key,
        limit=limit,
        sort_direction=sort_direction,
        state_filter=state_filter,
        control_filter=control_filter,
    )
    rows = connection.execute(sql, params).fetchall()
    return _apply_supplemental_net_price_fallback(
        connection,
        [_row_to_record(row, metric_key) for row in rows],
    )


def _query_peer_distribution(
    connection: sqlite3.Connection,
    focus_school: dict,
    metric_key: str,
) -> tuple[list[dict], str]:
    scope_parts = []

    state_filter = focus_school["state"] if focus_school["state"] else None
    control_codes = (focus_school["control_code"],) if focus_school["control_code"] else None

    if state_filter:
        scope_parts.append(focus_school['state'])
    if control_codes:
        scope_parts.append(focus_school["control"])

    sql, params = _build_distribution_sql(
        metric_key=metric_key,
        state_filter=state_filter,
        control_codes=control_codes,
    )
    rows = connection.execute(sql, params).fetchall()

    distribution = []
    normalized_focus = _normalize_lookup_value(focus_school["institution_name"])
    for row in rows:
        value = _coerce_float(row["metric_value"])
        if value is None:
            continue
        distribution.append(
            {
                "label": row["INSTNM"],
                "metric_value": value,
                "is_selected": _normalize_lookup_value(row["INSTNM"]) == normalized_focus,
            }
        )

    scope_text = " / ".join(scope_parts) if scope_parts else "all institutions"
    return distribution, scope_text


def _query_scoped_peer_distribution(
    connection: sqlite3.Connection,
    focus_school: dict,
    metric_key: str,
    scope_label: str,
    states: tuple[str, ...] | None = None,
    control_codes: tuple[str, ...] | None = None,
) -> tuple[list[dict], str]:
    sql, params = _build_distribution_sql(
        metric_key=metric_key,
        states=states,
        control_codes=control_codes,
    )
    rows = connection.execute(sql, params).fetchall()

    distribution = []
    normalized_focus = _normalize_lookup_value(focus_school["institution_name"])
    for row in rows:
        value = _coerce_float(row["metric_value"])
        if value is None:
            continue
        distribution.append(
            {
                "label": row["INSTNM"],
                "metric_value": value,
                "is_selected": _normalize_lookup_value(row["INSTNM"]) == normalized_focus,
            }
        )

    control_scope = _format_control_scope(control_codes)
    control_label = f"{control_scope.lower()} schools" if control_scope else "schools"

    return distribution, f"{scope_label} {control_label}"


def _query_benchmark_value(
    connection: sqlite3.Connection,
    metric_key: str,
    benchmark_stat: str,
    states: tuple[str, ...] | None = None,
    control_codes: tuple[str, ...] | None = None,
) -> float | None:
    sql, params = _build_benchmark_sql(
        metric_key=metric_key,
        benchmark_stat=benchmark_stat,
        states=states,
        control_codes=control_codes,
    )
    row = connection.execute(sql, params).fetchone()
    return _coerce_float(row[0]) if row else None


def _infer_focus_graph_scope(
    question: str,
    rows: list[dict],
    control_filter: tuple[str, ...] | None,
    matched_schools: list[str] | None = None,
) -> tuple[tuple[str, ...] | None, str | None]:
    lower_question = question.casefold()
    if _mentions_national_scope(question):
        return None, "Nationwide"
    if "dmv" in lower_question:
        return DMV_STATE_ABBRS, "DMV area"

    explicit_state = _extract_state_filter(question, ignore_phrases=matched_schools)
    if explicit_state:
        return (explicit_state,), explicit_state

    if control_filter:
        return None, "Nationwide"

    row_states = {row["state"] for row in rows if row.get("state")}
    if not control_filter and row_states and row_states.issubset(set(DMV_STATE_ABBRS)):
        return DMV_STATE_ABBRS, "DMV area"

    return None, None


def _average_row_field(rows: list[dict], field_name: str) -> float | None:
    values = [row[field_name] for row in rows if row.get(field_name) is not None]
    if not values:
        return None
    return float(sum(values) / len(values))


def _dedupe_preserve_order(values: list[str] | tuple[str, ...]) -> list[str]:
    ordered: list[str] = []
    seen = set()
    for value in values:
        if not value or value in seen:
            continue
        ordered.append(value)
        seen.add(value)
    return ordered


def _has_full_name_school_match(question: str, matched_schools: list[str]) -> bool:
    question_normalized = _normalize_lookup_value(question)
    return any(_normalize_lookup_value(school) in question_normalized for school in matched_schools)


def _extract_likely_school_phrase(question: str) -> str:
    match = re.search(
        r"\b([A-Z][A-Za-z&.'-]+(?:\s+[A-Z][A-Za-z&.'-]+){0,5}\s+(?:University|College|Institute|School))\b",
        question,
    )
    if match:
        return match.group(1).strip()

    uppercase_tokens = re.findall(r"\b[A-Z]{3,6}\b", question)
    if uppercase_tokens:
        return uppercase_tokens[0]

    return ""


def _build_unmatched_school_message(question: str) -> str | None:
    schools = _load_school_names()
    matched_schools = _match_schools(question, schools)
    metric_key, metric_detected, _ = _resolve_metric(question)
    state_filter = _extract_state_filter(question, ignore_phrases=matched_schools)
    control_filter = _extract_control_filter(question)
    wants_ranked_results = _wants_ranked_results(question)
    wants_distribution_graph = _wants_distribution_graph(question)
    has_context_metric = False

    if not _is_sql_answerable_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        has_context_metric=has_context_metric,
        context_schools=None,
    ):
        return None

    if matched_schools:
        return None

    likely_phrase = _extract_likely_school_phrase(question)
    if not likely_phrase:
        return None

    suggestions = difflib.get_close_matches(likely_phrase, schools, n=3, cutoff=0.65)
    if suggestions:
        suggestion_text = ", ".join(suggestions)
        return (
            f"I couldn't confidently match `{likely_phrase}` in the database. "
            f"Try one of these names instead: {suggestion_text}."
        )

    return f"I couldn't confidently match `{likely_phrase}` in the database. Try the institution's full official name."


def _build_no_data_message(
    question: str,
    context_schools: tuple[str, ...] | None = None,
    context_metric_key: str | None = None,
) -> str | None:
    matched_schools = _match_schools(question, _load_school_names())
    metric_key, metric_detected, _ = _resolve_metric(question)
    if not metric_detected and context_metric_key in METRIC_DEFINITIONS:
        metric_key = context_metric_key
    state_filter = _extract_state_filter(question, ignore_phrases=matched_schools)
    control_filter = _extract_control_filter(question)
    wants_ranked_results = _wants_ranked_results(question)
    wants_distribution_graph = _wants_distribution_graph(question)
    contextual_schools = list(context_schools or ())
    has_context_metric = bool(not metric_detected and context_metric_key in METRIC_DEFINITIONS)

    if wants_ranked_results and state_filter and matched_schools and not _has_full_name_school_match(question, matched_schools):
        matched_schools = []

    if contextual_schools and not matched_schools and _should_reuse_context_schools(
        question,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
    ):
        matched_schools = contextual_schools[:8]

    if not _is_sql_answerable_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        has_context_metric=has_context_metric,
        context_schools=tuple(contextual_schools),
    ):
        return None

    filters = []
    if state_filter:
        filters.append(STATE_ABBR_TO_NAME.get(state_filter, state_filter))
    if control_filter:
        control_scope = _format_control_scope(control_filter)
        if control_scope:
            filters.append(control_scope)

    school_context = matched_schools or contextual_schools
    if school_context:
        return (
            f"I couldn't find rows in the current database for "
            f"{_format_school_series(school_context[:3])} on {METRIC_DEFINITIONS[metric_key]['short_label'].lower()}."
        )

    if filters:
        return (
            f"I couldn't find rows in the current database for "
            f"{METRIC_DEFINITIONS[metric_key]['short_label'].lower()} with filters {', '.join(filters)}."
        )

    return None


def _build_sql_fallback_instruction(
    unmatched_school_message: str | None,
    no_data_message: str | None,
) -> str:
    repository_gap = unmatched_school_message or no_data_message or (
        "The local SQLite repository did not return a usable school-level result for this prompt."
    )
    return (
        "The user asked for school-level structured data, but the local SQLite repository could not satisfy the request. "
        f"Repository limitation: {repository_gap} "
        "Do not imply that you searched the web or retrieved live external data. "
        "Start with one brief sentence explaining that the requested specific data is not available in the current repository, "
        "then give the best concise general guidance or background you can from model knowledge."
    )


def _get_recent_sql_payload(chat_history: list[dict] | None) -> dict | None:
    if not chat_history:
        return None

    for entry in reversed(chat_history):
        if entry.get("role") == "assistant" and entry.get("kind") == "sql":
            payload = entry.get("payload")
            if isinstance(payload, dict):
                return payload
    return None


def _extract_context_schools_from_payload(payload: dict | None) -> list[str]:
    if not payload:
        return []

    schools = payload.get("matched_schools") or [
        row.get("institution_name")
        for row in payload.get("rows", [])
        if row.get("institution_name")
    ]
    return _dedupe_preserve_order(schools)[:8]


def _should_reuse_school_context(
    question: str,
    prior_sql_payload: dict | None,
    matched_schools: list[str],
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
) -> bool:
    if not prior_sql_payload or matched_schools or state_filter or control_filter:
        return False

    if _looks_like_model_preferred_question(question):
        return False

    lower_question = question.casefold()
    prior_result_kind = prior_sql_payload.get("result_kind")
    if prior_result_kind == "ranking":
        return any(
            re.search(rf"\b{re.escape(phrase)}\b", lower_question)
            for phrase in DIRECT_CONTEXT_REFERENCE_PHRASES
        )

    return any(phrase in lower_question for phrase in FOLLOW_UP_SCHOOL_PHRASES)


def _should_reuse_metric_context(
    question: str,
    prior_sql_payload: dict | None,
    metric_detected: bool,
    matched_schools: list[str],
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
) -> bool:
    if not prior_sql_payload or metric_detected:
        return False

    if _looks_like_model_preferred_question(question):
        return False

    lower_question = question.casefold()
    is_follow_up = any(phrase in lower_question for phrase in FOLLOW_UP_SCHOOL_PHRASES)
    if not is_follow_up:
        return False

    return bool(
        state_filter
        or control_filter
        or wants_ranked_results
        or wants_distribution_graph
    )


def _should_reuse_context_schools(
    question: str,
    metric_detected: bool,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
) -> bool:
    return bool(
        metric_detected
        or wants_ranked_results
        or wants_distribution_graph
        or state_filter
        or control_filter
    )


def _looks_like_model_preferred_question(question: str) -> bool:
    lower_question = question.casefold()
    return any(phrase in lower_question for phrase in MODEL_PREFERRED_INTENT_PHRASES)


def _has_resume_context_available() -> bool:
    return bool(str(st.session_state.get("resume_text", "")).strip())


def _has_applicant_stats_available() -> bool:
    return any(
        str(st.session_state.get(field_name, "")).strip()
        for field_name in (
            "app_info_gpa",
            "app_info_sat",
            "app_info_act",
            "app_info_home_state",
            "app_info_major_interest",
        )
    )


def _looks_like_personalized_advice_question(question: str) -> bool:
    lower_question = question.casefold()

    if any(phrase in lower_question for phrase in MODEL_PREFERRED_INTENT_PHRASES):
        return True
    if any(phrase in lower_question for phrase in PERSONALIZED_ADVICE_INTENT_PHRASES):
        return True
    if FIRST_PERSON_REFERENCE_PATTERN.search(question) and any(
        token in lower_question for token in APPLICANT_STAT_QUERY_TOKENS
    ):
        return True
    return False


def _has_school_data_scope(
    matched_schools: list[str],
    context_schools: tuple[str, ...] | None,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
) -> bool:
    return bool(matched_schools or context_schools or state_filter or control_filter)


def _looks_like_school_stats_question(
    question: str,
    matched_schools: list[str],
    metric_detected: bool,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
    context_schools: tuple[str, ...] | None = None,
) -> bool:
    lower_question = question.casefold()
    has_scope = _has_school_data_scope(
        matched_schools=matched_schools,
        context_schools=context_schools,
        state_filter=state_filter,
        control_filter=control_filter,
    )
    has_profile_intent = any(phrase in lower_question for phrase in DATA_PROFILE_INTENT_PHRASES)
    has_comparison_intent = any(phrase in lower_question for phrase in DATA_COMPARISON_INTENT_PHRASES)
    has_profile_request = has_profile_intent and bool(matched_schools or context_schools)
    has_data_shape = bool(
        metric_detected
        or wants_distribution_graph
        or has_profile_request
        or (has_comparison_intent and has_scope)
    )

    if not has_data_shape:
        return False
    if not has_scope and not wants_ranked_results and not wants_distribution_graph:
        return False

    return True


def _classify_intent(
    question: str,
    matched_schools: list[str],
    metric_detected: bool,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
    context_schools: tuple[str, ...] | None = None,
) -> str:
    if _looks_like_personalized_advice_question(question):
        return INTENT_PERSONALIZED_ADVICE
    if looks_like_rag_question(question):
        return INTENT_MAJOR_OR_CAREER
    if _looks_like_school_stats_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        context_schools=context_schools,
    ):
        return INTENT_SCHOOL_STATS
    return INTENT_GENERAL


def _build_source_plan(
    question: str,
    matched_schools: list[str],
    metric_detected: bool,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
    context_schools: tuple[str, ...] | None = None,
) -> dict:
    intent = _classify_intent(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        context_schools=context_schools,
    )
    school_stats_signal = _looks_like_school_stats_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        context_schools=context_schools,
    )
    use_resume = intent == INTENT_PERSONALIZED_ADVICE and _has_resume_context_available()
    use_applicant_stats = (
        intent == INTENT_PERSONALIZED_ADVICE and _has_applicant_stats_available()
    )

    return {
        "intent": intent,
        "use_resume": use_resume,
        "use_applicant_stats": use_applicant_stats,
        "use_rag": looks_like_rag_question(question),
        "use_sql_result": intent == INTENT_SCHOOL_STATS,
        "use_sql_context": intent == INTENT_PERSONALIZED_ADVICE and school_stats_signal,
    }


def _is_sql_answerable_question(
    question: str,
    matched_schools: list[str],
    metric_detected: bool,
    wants_ranked_results: bool,
    wants_distribution_graph: bool,
    state_filter: str | None,
    control_filter: tuple[str, ...] | None,
    has_context_metric: bool = False,
    context_schools: tuple[str, ...] | None = None,
) -> bool:
    return _looks_like_school_stats_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected or has_context_metric,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        context_schools=context_schools,
    )


def _build_summary_text(
    result_kind: str,
    metric_key: str,
    rows: list[dict],
    benchmark: dict | None = None,
    requested_metrics: list[str] | None = None,
) -> str:
    if not rows:
        return ""

    metric_label = METRIC_DEFINITIONS[metric_key]["short_label"].lower()
    lower_is_better = METRIC_DEFINITIONS[metric_key]["default_direction"] == "ASC"
    valid_rows = [row for row in rows if row.get("metric_value") is not None]
    if not valid_rows:
        valid_rows = rows

    cheapest_row = min(
        (row for row in rows if row.get("net_price") is not None),
        key=lambda row: row["net_price"],
        default=None,
    )
    strongest_earnings_row = max(
        (row for row in rows if row.get("earnings") is not None),
        key=lambda row: row["earnings"],
        default=None,
    )
    lowest_debt_row = min(
        (row for row in rows if row.get("debt") is not None),
        key=lambda row: row["debt"],
        default=None,
    )

    if result_kind == "school_profile":
        school = rows[0]
        requested_metrics = requested_metrics or [metric_key]
        if benchmark and school.get("metric_value") is not None and benchmark.get("value") is not None:
            difference_value = school["metric_value"] - benchmark["value"]
            if abs(difference_value) < 0.01:
                summary = (
                    f"{school['institution_name']} comes in at "
                    f"{_format_metric_value(metric_key, school['metric_value'])} for {metric_label}, "
                    f"which is right in line with the {benchmark['label'].lower()} "
                    f"({_format_metric_value(metric_key, benchmark['value'])})."
                )
            else:
                direction = "above" if difference_value > 0 else "below"
                summary = (
                    f"{school['institution_name']} comes in at "
                    f"{_format_metric_value(metric_key, school['metric_value'])} for {metric_label}, "
                    f"{_format_metric_value(metric_key, abs(difference_value))} {direction} the "
                    f"{benchmark['label'].lower()} ({_format_metric_value(metric_key, benchmark['value'])})."
                )
        else:
            summary = (
                f"{school['institution_name']} comes in at "
                f"{_format_metric_value(metric_key, school['metric_value'])} for {metric_label}."
            )

        additional_requested = [key for key in requested_metrics if key != metric_key]
        if additional_requested:
            detail_bits = []
            for extra_metric in additional_requested[:2]:
                if school.get(extra_metric) is not None:
                    detail_bits.append(
                        f"{METRIC_DEFINITIONS[extra_metric]['short_label'].lower()} "
                        f"of {_format_metric_value(extra_metric, school[extra_metric])}"
                    )
            if detail_bits:
                summary += " It also shows " + " and ".join(detail_bits) + "."
                return _compact_display_text(
                    summary,
                    max_words=MAX_SQL_SUMMARY_WORDS,
                    max_sentences=MAX_SQL_SUMMARY_SENTENCES,
                    max_paragraphs=1,
                )

        if metric_key == "earnings" and school.get("net_price") is not None and school.get("debt") is not None:
            summary += (
                f" For the decision, weigh that against net price of {_format_currency(school['net_price'])} "
                f"and median debt of {_format_currency(school['debt'])}."
            )
        elif metric_key in {"net_price", "tuition_in", "tuition_out", "debt"} and school.get("earnings") is not None:
            summary += (
                f" For the decision, weigh that against median earnings of {_format_currency(school['earnings'])}."
            )
        elif metric_key == "admission_rate" and school.get("net_price") is not None and school.get("earnings") is not None:
            summary += (
                f" For the decision, balance that access picture with net price of {_format_currency(school['net_price'])} "
                f"and median earnings of {_format_currency(school['earnings'])}."
            )
        return _compact_display_text(
            summary,
            max_words=MAX_SQL_SUMMARY_WORDS,
            max_sentences=MAX_SQL_SUMMARY_SENTENCES,
            max_paragraphs=1,
        )

    ranked_rows = sorted(
        valid_rows,
        key=lambda row: row["metric_value"] if row.get("metric_value") is not None else float("inf"),
        reverse=not lower_is_better,
    )
    best_row = ranked_rows[0]
    edge_word = "lowest" if lower_is_better else "highest"

    if metric_key == "admission_rate" and len(ranked_rows) >= 2:
        comparison_row = min(ranked_rows, key=lambda row: row["metric_value"] if row.get("metric_value") is not None else 1)
        summary = (
            f"{best_row['institution_name']} is the most accessible option here at "
            f"{_format_metric_value(metric_key, best_row['metric_value'])}, while {comparison_row['institution_name']} "
            f"is the most selective at {_format_metric_value(metric_key, comparison_row['metric_value'])}."
        )
    else:
        summary = (
            f"{best_row['institution_name']} stands out on {metric_label} at "
            f"{_format_metric_value(metric_key, best_row['metric_value'])}, which is the {edge_word} value in this set."
        )

    if cheapest_row and strongest_earnings_row and cheapest_row["institution_name"] != strongest_earnings_row["institution_name"]:
        summary += (
            f" If you're deciding between these options, the clearest tradeoff is "
            f"{strongest_earnings_row['institution_name']} for stronger earnings "
            f"versus {cheapest_row['institution_name']} for lower net price."
        )
    elif lowest_debt_row and best_row["institution_name"] != lowest_debt_row["institution_name"] and metric_key != "debt":
        summary += (
            f" If you're deciding between these options, compare that with {lowest_debt_row['institution_name']}'s "
            f"lower median debt of {_format_currency(lowest_debt_row['debt'])}."
        )
    elif result_kind == "ranking":
        summary += " Use net price and debt to decide whether the top option also looks strongest overall."

    return _compact_display_text(
        summary,
        max_words=MAX_SQL_SUMMARY_WORDS,
        max_sentences=MAX_SQL_SUMMARY_SENTENCES,
        max_paragraphs=1,
    )


def _append_chat_history_entry(entry: dict) -> None:
    history = st.session_state.setdefault("chat_history", [])
    history.append(entry)


def _split_into_sentences(text: str) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []
    return [part.strip() for part in re.split(r"(?<=[.!?])\s+", cleaned) if part.strip()]


def _history_to_model_messages(chat_history: list[dict]) -> list[dict]:
    if not chat_history:
        return []

    conversation_entries = [
        entry
        for entry in chat_history
        if entry.get("role") in {"user", "assistant"}
    ]
    recent_entries = conversation_entries[-(MAX_HISTORY_TURNS_FOR_MODEL * 2) :]

    messages = []
    for entry in recent_entries:
        if entry["role"] == "user":
            messages.append({"role": "user", "content": entry["content"]})
            continue

        if entry.get("kind") == "sql":
            messages.append(
                {
                    "role": "assistant",
                    "content": (
                        "Structured data result shown to the user:\n"
                        f"{entry.get('summary_text', '')}"
                    ).strip(),
                }
            )
        else:
            messages.append({"role": "assistant", "content": entry["content"]})

    return messages


def _history_to_transcript(chat_history: list[dict]) -> str:
    lines = []
    for entry in chat_history:
        if entry["role"] == "user":
            lines.append(f"User: {entry['content']}")
            continue

        if entry.get("kind") == "sql":
            payload = entry["payload"]
            lines.append(f"Assistant (SQL): {entry.get('summary_text', '')}")
            for row in payload.get("rows", [])[:5]:
                metric_key = payload["metric_key"]
                metric_label = METRIC_DEFINITIONS[metric_key]["short_label"]
                lines.append(
                    f"- {row['institution_name']} | {metric_label} {_format_metric_value(metric_key, row['metric_value'])} | "
                    f"net price {_format_currency(row['net_price'])} | debt {_format_currency(row['debt'])} | "
                    f"admission {_format_percent(row['admission_rate'])} | completion {_format_percent(row['completion_rate'])}"
                )
        else:
            lines.append(f"Assistant: {entry['content']}")
    return "\n".join(lines)


def _ensure_terminal_period(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", str(text).strip())
    if not cleaned:
        return ""
    if cleaned[-1] not in ".!?":
        cleaned += "."
    return cleaned


def _format_school_series(schools: list[str]) -> str:
    items = [school for school in schools if school]
    if not items:
        return "the schools discussed"
    if len(items) == 1:
        return items[0]
    if len(items) == 2:
        return f"{items[0]} and {items[1]}"
    return f"{', '.join(items[:-1])}, and {items[-1]}"


def _format_text_block_html(text: str, class_name: str) -> str:
    escaped = html.escape(str(text).strip())
    body = escaped.replace("\n", "<br>") if escaped else ""
    return f"<div class='{class_name}'>{body}</div>"


def _render_assistant_text(text: str, placeholder=None) -> None:
    target = placeholder if placeholder is not None else st
    target.markdown(
        _format_text_block_html(text, "edu-chat-copy"),
        unsafe_allow_html=True,
    )


def _render_streamed_assistant_reply(deployment_id: str, messages: list[dict]) -> str:
    placeholder = st.empty()
    chunks: list[str] = []
    for chunk in _stream_chat_with_model(deployment_id, messages):
        chunk_text = str(chunk)
        if not chunk_text:
            continue
        chunks.append(chunk_text)
        _render_assistant_text("".join(chunks), placeholder=placeholder)

    reply = "".join(chunks).strip()
    if reply:
        _render_assistant_text(reply, placeholder=placeholder)
    return reply


def _normalize_summary_items(
    proposed_items: list[str],
    fallback_items: list[str],
    validator,
) -> list[str]:
    normalized: list[str] = []
    seen = set()

    for source in (proposed_items, fallback_items):
        for item in source:
            cleaned = _ensure_terminal_period(item)
            if not cleaned or cleaned in seen or not validator(cleaned):
                continue
            normalized.append(cleaned)
            seen.add(cleaned)
            if len(normalized) == 3:
                return normalized

    return normalized


def _is_end_session_takeaway(text: str) -> bool:
    cleaned = _ensure_terminal_period(text)
    if not cleaned or len(cleaned) > 240:
        return False

    lower_text = cleaned.casefold()
    if any(blocked_phrase in lower_text for blocked_phrase in END_SESSION_BLOCKED_PHRASES):
        return False

    if cleaned.count("$") > 0 or cleaned.count("%") > 0:
        return False

    if re.search(r"\b(gpa|sat|act)\b", lower_text) and any(character.isdigit() for character in cleaned):
        return False

    return len(cleaned.split()) >= 8


def _is_end_session_next_step(text: str) -> bool:
    cleaned = _ensure_terminal_period(text)
    if not cleaned or len(cleaned) > 220:
        return False

    lower_text = cleaned.casefold()
    if any(blocked_phrase in lower_text for blocked_phrase in END_SESSION_BLOCKED_PHRASES):
        return False

    first_word = re.sub(r"[^A-Za-z].*", "", cleaned.split()[0]).capitalize()
    if first_word not in END_SESSION_NEXT_STEP_VERBS:
        return False

    return len(cleaned.split()) >= 6


def _is_end_session_closing_thought(text: str) -> bool:
    cleaned = _ensure_terminal_period(text)
    if not cleaned or len(cleaned) > 180:
        return False

    lower_text = cleaned.casefold()
    if any(blocked_phrase in lower_text for blocked_phrase in END_SESSION_BLOCKED_PHRASES):
        return False

    return len(cleaned.split()) >= 6


def _normalize_closing_thought(proposed_text: str, fallback_text: str) -> str:
    for candidate in (proposed_text, fallback_text):
        cleaned = _ensure_terminal_period(candidate)
        if cleaned and _is_end_session_closing_thought(cleaned):
            return cleaned
    return "You have more clarity now, and the next step should feel more manageable."


def _build_end_session_takeaway_candidates(
    discussed_schools: list[str],
    discussed_metrics: set[str],
    major_interest: str,
) -> list[str]:
    school_phrase = _format_school_series(discussed_schools[:3])
    has_cost_focus = bool({"net_price", "debt", "tuition_in", "tuition_out"} & discussed_metrics)
    has_outcome_focus = bool({"earnings", "completion_rate"} & discussed_metrics)
    has_selectivity_focus = "admission_rate" in discussed_metrics
    candidates: list[str] = []

    if len(discussed_schools) >= 2:
        candidates.append(
            f"You seem to make clearer decisions when schools are compared side by side, which is helping your shortlist become more grounded and deliberate."
        )
    elif discussed_schools:
        candidates.append(
            f"{discussed_schools[0]} appears to be acting as a reference point, which can help clarify what matters most as you compare other options."
        )

    if has_cost_focus and has_outcome_focus:
        candidates.append(
            "There is a clear pattern of balancing affordability with long-term outcomes, which is a strong way to judge overall fit."
        )
    elif has_cost_focus:
        candidates.append(
            "Cost clearly matters in your decision-making, which should help you build a list that feels realistic as well as appealing."
        )
    elif has_outcome_focus:
        candidates.append(
            "You are leaning toward practical outcomes, which suggests you want options that feel worthwhile after graduation and not just attractive on paper."
        )

    if has_selectivity_focus:
        candidates.append(
            "You seem to care about balancing ambition with realism, which is useful for building a healthier mix of reach, target, and likely schools."
        )

    if major_interest:
        candidates.append(
            f"Your interest in {major_interest} can serve as a strong organizing theme for narrowing schools and evaluating which programs best support your direction."
        )

    if len(discussed_schools) >= 2 and not has_cost_focus and not has_outcome_focus:
        candidates.append(
            f"You are moving beyond a general impression of {school_phrase} and toward a more structured sense of fit, which should make future choices easier."
        )

    candidates.extend(
        [
            "You are starting to evaluate colleges through tradeoffs rather than hype alone, which is usually where better decisions begin to take shape.",
            "The conversation points toward a clearer sense of what matters most to you, even if the final shortlist is still evolving.",
            "Your questions suggest you are looking for a college path that feels both personally fitting and practically worthwhile."
        ]
    )
    return candidates


def _build_end_session_next_step_candidates(
    discussed_schools: list[str],
    discussed_metrics: set[str],
    major_interest: str,
) -> list[str]:
    school_phrase = _format_school_series(discussed_schools[:3])
    candidates: list[str] = []

    if len(discussed_schools) >= 2:
        candidates.append(
            f"Build a shortlist from {school_phrase} and compare those options in one place on fit, cost, and long-term value."
        )
    else:
        candidates.append(
            "Build a shortlist of three to five schools that still feel strongest so you can compare them more intentionally."
        )

    if major_interest:
        candidates.append(
            f"Explore how {major_interest} is structured at your top schools, including flexibility, internships, and likely career paths."
        )

    if not {"net_price", "debt", "tuition_in", "tuition_out"} & discussed_metrics:
        candidates.append(
            "Compare likely cost and aid across your top options so affordability becomes part of the decision before your list grows."
        )
    else:
        candidates.append(
            "Review financial aid, net price, and borrowing together for your leading schools so the cost picture stays realistic."
        )

    if "admission_rate" not in discussed_metrics:
        candidates.append(
            "Refine your list into reach, target, and likely schools so your application strategy stays balanced."
        )
    else:
        candidates.append(
            "Check whether the schools you still like most also make sense within a balanced reach, target, and likely mix."
        )

    if not {"earnings", "completion_rate"} & discussed_metrics:
        candidates.append(
            "Research outcomes and support systems for the majors you are considering so your shortlist reflects more than name recognition."
        )
    else:
        candidates.append(
            "Ask whether your leading schools still feel strongest once you weigh outcomes alongside cost and day-to-day fit."
        )

    candidates.append(
        "Reach out to a current student, admissions counselor, or program contact at one top-choice school to pressure-test what it would actually feel like to attend."
    )
    return candidates


def _default_end_session_closing_thought(
    discussed_schools: list[str],
    major_interest: str,
) -> str:
    if major_interest:
        return (
            f"You are getting closer to a clearer direction, and that should make your next round of {major_interest} and school comparisons feel much more focused."
        )
    if discussed_schools:
        return (
            "You have enough clarity now to make the next round of school decisions feel more focused and less overwhelming."
        )
    return "You are asking the right questions, and that clarity should make the next step in your college search feel much more manageable."


def _extract_json_object_from_text(text: str) -> dict:
    raw_text = str(text or "").strip()
    if not raw_text:
        return {}

    fence_match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", raw_text, re.DOTALL)
    if fence_match:
        raw_text = fence_match.group(1).strip()

    try:
        return json.loads(raw_text)
    except json.JSONDecodeError:
        pass

    start = raw_text.find("{")
    end = raw_text.rfind("}")
    if start == -1 or end == -1 or end <= start:
        return {}

    try:
        return json.loads(raw_text[start : end + 1])
    except json.JSONDecodeError:
        return {}


def _heuristic_session_summary(chat_history: list[dict]) -> dict:
    sql_payloads = [
        entry["payload"]
        for entry in chat_history
        if entry.get("role") == "assistant" and entry.get("kind") == "sql" and entry.get("payload")
    ]
    discussed_metrics = {payload["metric_key"] for payload in sql_payloads}
    discussed_schools = _dedupe_preserve_order(
        [
            school
            for payload in sql_payloads
            for school in (
                payload.get("matched_schools")
                or [row.get("institution_name") for row in payload.get("rows", []) if row.get("institution_name")]
            )
        ]
    )
    major_interest = str(st.session_state.get("app_info_major_interest", "")).strip()

    takeaways = _normalize_summary_items(
        _build_end_session_takeaway_candidates(
            discussed_schools=discussed_schools,
            discussed_metrics=discussed_metrics,
            major_interest=major_interest,
        ),
        fallback_items=[
            "You are building a clearer sense of what matters most in your college decision.",
            "The conversation is helping narrow your options through fit, practicality, and overall direction.",
            "Your questions suggest you want a college choice that feels thoughtful, realistic, and aligned with your goals.",
        ],
        validator=_is_end_session_takeaway,
    )

    next_steps = _normalize_summary_items(
        _build_end_session_next_step_candidates(
            discussed_schools=discussed_schools,
            discussed_metrics=discussed_metrics,
            major_interest=major_interest,
        ),
        fallback_items=[
            "Build a shortlist of the schools that still feel strongest and compare them in one place.",
            "Compare cost, outcomes, and overall fit before deciding which schools deserve more attention.",
            "Refine your list into a realistic mix so the next stage of the process feels more manageable.",
        ],
        validator=_is_end_session_next_step,
    )

    return {
        "takeaways": takeaways[:3],
        "next_steps": next_steps[:3],
        "closing_thought": _default_end_session_closing_thought(
            discussed_schools=discussed_schools,
            major_interest=major_interest,
        ),
        "source": "heuristic",
    }


def _generate_session_summary(chat_history: list[dict]) -> dict:
    if not chat_history:
        return {
            "takeaways": [],
            "next_steps": [],
            "closing_thought": "",
            "source": "empty",
        }

    transcript = _history_to_transcript(chat_history)
    heuristic_summary = _heuristic_session_summary(chat_history)
    applicant_context = _build_application_information_context()
    resume_text = str(st.session_state.get("resume_text", "")).strip()

    openai.api_type = "azure"
    openai.api_base = os.getenv("AZURE_OPENAI_ENDPOINT")
    openai.api_version = os.getenv("AZURE_OPENAI_API_VERSION", "2025-04-01-preview")
    api_key = os.getenv("AZURE_OPENAI_API_KEY")
    deployment_id = os.getenv("AZURE_OPENAI_DEPLOYMENT_ID")

    if api_key and deployment_id and openai.api_base:
        try:
            openai.api_key = api_key
            context_blocks = [f"Conversation transcript:\n{_truncate_text(transcript, 14000)}"]
            if applicant_context:
                context_blocks.append(applicant_context)
            if resume_text:
                context_blocks.append(
                    "Resume context provided by the student:\n"
                    + _truncate_text(resume_text, 3000)
                )
            messages = [
                {
                    "role": "system",
                    "content": END_SESSION_SUMMARY_SYSTEM_MESSAGE,
                },
                {
                    "role": "user",
                    "content": "\n\n".join(context_blocks),
                },
            ]
            reply = _chat_with_model(deployment_id, messages)
            parsed = _extract_json_object_from_text(reply)
            takeaways = _normalize_summary_items(
                [str(item).strip() for item in parsed.get("top_takeaways", []) if str(item).strip()],
                fallback_items=heuristic_summary["takeaways"],
                validator=_is_end_session_takeaway,
            )
            next_steps = _normalize_summary_items(
                [str(item).strip() for item in parsed.get("next_steps", []) if str(item).strip()],
                fallback_items=heuristic_summary["next_steps"],
                validator=_is_end_session_next_step,
            )
            closing_thought = _normalize_closing_thought(
                str(parsed.get("closing_thought", "")).strip(),
                heuristic_summary["closing_thought"],
            )
            if len(takeaways) == 3 and len(next_steps) == 3 and closing_thought:
                return {
                    "takeaways": takeaways,
                    "next_steps": next_steps,
                    "closing_thought": closing_thought,
                    "source": "model",
                }
        except Exception:
            logger.exception("Session summary generation fell back to heuristic mode.")

    return _heuristic_session_summary(chat_history)


def _build_summary_stats(
    result_kind: str,
    rows: list[dict],
    metric_key: str | None = None,
    benchmark: dict | None = None,
) -> list[dict]:
    if not rows:
        return []

    if benchmark and metric_key:
        focus_value = rows[0]["metric_value"]
        benchmark_value = benchmark.get("value")
        difference_value = (
            focus_value - benchmark_value
            if focus_value is not None and benchmark_value is not None
            else None
        )
        focus_label = f"School {METRIC_DEFINITIONS[metric_key]['short_label'].lower()}"
        return [
            {
                "label": focus_label,
                "value": _format_metric_value(metric_key, focus_value),
            },
            {
                "label": benchmark["label"],
                "value": _format_metric_value(metric_key, benchmark_value),
            },
            {
                "label": "Difference",
                "value": _format_metric_value(metric_key, difference_value),
            },
        ]

    if result_kind == "school_profile":
        row = rows[0]
        return [
            {"label": "Median earnings", "value": _format_currency(row["earnings"])},
            {
                "label": METRIC_DEFINITIONS["net_price"]["label"],
                "value": _format_currency(row["net_price"]),
            },
            {"label": "Median debt", "value": _format_currency(row["debt"])},
        ]

    return [
        {
            "label": "Average median earnings",
            "value": _format_currency(_average_row_field(rows, "earnings")),
        },
        {
            "label": METRIC_DEFINITIONS["net_price"]["label"],
            "value": _format_currency(_average_row_field(rows, "net_price")),
        },
        {
            "label": "Average median debt",
            "value": _format_currency(_average_row_field(rows, "debt")),
        },
    ]


def _build_display_table(metric_key: str, rows: list[dict]) -> pd.DataFrame:
    metric_label = METRIC_DEFINITIONS[metric_key]["label"]
    return pd.DataFrame(
        [
            {
                "School": row["institution_name"],
                "Location": row["location"],
                "Type of University": row["control"],
                metric_label: _format_metric_value(metric_key, row["metric_value"]),
                METRIC_DEFINITIONS["net_price"]["label"]: _format_currency(row["net_price"]),
                "Median debt": _format_currency(row["debt"]),
                "Admission rate": _format_percent(row["admission_rate"]),
                "Undergraduates": _format_count(row["undergrad_size"]),
                "Completion rate": _format_percent(row["completion_rate"]),
            }
            for row in rows
        ]
    )


def _build_structured_context_from_payload(payload: dict) -> str:
    metric_key = payload["metric_key"]
    metric_label = METRIC_DEFINITIONS[metric_key]["label"]
    lines = [
        "SQLite structured context",
        "Dataset: Most Recent Cohorts Institution (institution-level College Scorecard slice).",
        f"Primary metric: {metric_label}",
    ]

    if payload["matched_schools"]:
        lines.append("Matched schools: " + ", ".join(payload["matched_schools"]))
    if payload["state_filter"]:
        lines.append(f"State filter: {payload['state_filter']}")
    if payload["control_filter"]:
        control_scope = _format_control_scope(payload["control_filter"])
        if control_scope:
            lines.append("Control filter: " + control_scope)
    if payload["assumptions"]:
        lines.append("Assumptions: " + " ".join(payload["assumptions"]))

    lines.append("Returned rows:")
    for row in payload["rows"]:
        net_price_text = _format_currency(row["net_price"])
        if row.get("net_price_source") == "supplemental_history":
            net_price_text += " (supplemental history source)"
        lines.append(
            f"- {row['institution_name']} | {row['control']} | {row['location']} | "
            f"{metric_label} {_format_metric_value(metric_key, row['metric_value'])} | "
            f"net price {net_price_text} | "
            f"median debt {_format_currency(row['debt'])} | "
            f"admission {_format_percent(row['admission_rate'])} | "
            f"completion {_format_percent(row['completion_rate'])}"
        )

    if payload["peer_distribution"]:
        values = [item["metric_value"] for item in payload["peer_distribution"]]
        if values:
            lines.append(
                f"Peer range for focus school ({payload['peer_scope']}): "
                f"{_format_metric_value(metric_key, min(values))} to "
                f"{_format_metric_value(metric_key, max(values))} across {len(values)} schools."
            )
    if payload.get("benchmark"):
        lines.append(
            f"Benchmark used: {payload['benchmark']['label']} = "
            f"{_format_metric_value(metric_key, payload['benchmark'].get('value'))}."
        )

    lines.append(
        "Important limitation: this database is school-level only, so it should not be used for major-specific claims."
    )
    if payload.get("net_price_history") and payload.get("result_kind") == "school_profile":
        lines.append(
            "Supplemental net price history is available for recent academic years and can be broken out by family income bracket."
        )
    return "\n".join(lines)


def _build_sql_query_display_items(payload: dict) -> list[dict]:
    items = payload.get("sql_queries", [])
    formatted_items = []
    for item in items:
        sql = item.get("sql", "").strip()
        if not sql:
            continue
        formatted_items.append(
            {
                "label": item.get("label", "SQL query"),
                "sql": sql,
            }
        )
    return formatted_items


def _build_sql_result_payload(
    question: str,
    context_schools: tuple[str, ...] | None = None,
    context_metric_key: str | None = None,
) -> dict | None:
    if not question.strip():
        return None

    if not SQLITE_PATH.exists() and not DATASET_PATH.exists():
        return None

    _ensure_sqlite_database()

    schools = _load_school_names()
    matched_schools = _match_schools(question, schools)
    metric_key, metric_detected, assumptions = _resolve_metric(question)
    metric_mentions = _extract_metric_mentions(question)
    has_context_metric = bool(not metric_detected and context_metric_key in METRIC_DEFINITIONS)
    if not metric_detected and context_metric_key in METRIC_DEFINITIONS:
        metric_key = context_metric_key
        assumptions.append(
            "Using the previously discussed metric for this follow-up."
        )
    requested_metrics = metric_mentions or [metric_key]
    state_filter = _extract_state_filter(question, ignore_phrases=matched_schools)
    control_filter = _extract_control_filter(question)
    wants_ranked_results = _wants_ranked_results(question)
    wants_distribution_graph = _wants_distribution_graph(question)
    benchmark_stat = _extract_benchmark_stat(question)
    limit = _extract_requested_limit(question)
    sort_direction = _extract_sort_direction(question, metric_key)

    contextual_schools = _dedupe_preserve_order(list(context_schools or ()))
    lower_question = question.casefold()
    if wants_ranked_results and state_filter and matched_schools and not _has_full_name_school_match(question, matched_schools):
        matched_schools = []

    if matched_schools and contextual_schools and any(
        token in lower_question for token in ["compare", "versus", " vs ", "against", "relative to"]
    ):
        merged_schools = _dedupe_preserve_order(contextual_schools + matched_schools)
        if len(merged_schools) > len(matched_schools):
            matched_schools = merged_schools[:8]
            assumptions.append(
                "Using the previously discussed school context to complete this comparison."
            )

    if (
        not matched_schools
        and contextual_schools
        and _should_reuse_context_schools(
            question,
            metric_detected=metric_detected,
            wants_ranked_results=wants_ranked_results,
            wants_distribution_graph=wants_distribution_graph,
            state_filter=state_filter,
            control_filter=control_filter,
        )
    ):
        matched_schools = contextual_schools[:8]
        assumptions.append(
            "Using the previously discussed school context for this follow-up."
        )
        state_filter = _extract_state_filter(question, ignore_phrases=matched_schools)

    if _looks_like_unmatched_school_request(
        question,
        matched_schools=matched_schools,
        wants_ranked_results=wants_ranked_results,
        state_filter=state_filter,
        control_filter=control_filter,
    ):
        return None

    if not _is_sql_answerable_question(
        question,
        matched_schools=matched_schools,
        metric_detected=metric_detected,
        wants_ranked_results=wants_ranked_results,
        wants_distribution_graph=wants_distribution_graph,
        state_filter=state_filter,
        control_filter=control_filter,
        has_context_metric=has_context_metric,
        context_schools=tuple(contextual_schools),
    ):
        return None

    with sqlite3.connect(SQLITE_PATH) as connection:
        connection.row_factory = sqlite3.Row
        benchmark = None
        sql_queries: list[dict] = []
        net_price_history: list[dict] = []

        preferred_graph = "comparison"
        auto_show_graph = False

        if matched_schools:
            primary_sql, primary_params = _build_fetch_institutions_sql(matched_schools, metric_key)
            rows = _fetch_institutions_by_name(connection, matched_schools, metric_key)
            if not rows:
                return None
            sql_queries.append(
                {
                    "label": "Primary result query",
                    "sql": _render_sql_with_params(primary_sql, primary_params),
                }
            )
            result_kind = "school_profile" if len(rows) == 1 else "comparison"
            focus_row = rows[0]
            scoped_states, scoped_scope_label = _infer_focus_graph_scope(
                question,
                rows,
                control_filter,
                matched_schools=matched_schools,
            )
            if wants_distribution_graph and scoped_scope_label and focus_row:
                scoped_control_codes = control_filter or (
                    (focus_row["control_code"],) if focus_row.get("control_code") else None
                )
                distribution_sql, distribution_params = _build_distribution_sql(
                    metric_key=metric_key,
                    states=scoped_states,
                    control_codes=scoped_control_codes,
                )
                peer_distribution, peer_scope = _query_scoped_peer_distribution(
                    connection,
                    focus_school=focus_row,
                    metric_key=metric_key,
                    scope_label=scoped_scope_label,
                    states=scoped_states,
                    control_codes=scoped_control_codes,
                )
                sql_queries.append(
                    {
                        "label": "Graph distribution query",
                        "sql": _render_sql_with_params(distribution_sql, distribution_params),
                    }
                )
                benchmark_sql, benchmark_params = _build_benchmark_sql(
                    metric_key=metric_key,
                    benchmark_stat=benchmark_stat,
                    states=scoped_states,
                    control_codes=scoped_control_codes,
                )
                benchmark_value = _query_benchmark_value(
                    connection,
                    metric_key=metric_key,
                    benchmark_stat=benchmark_stat,
                    states=scoped_states,
                    control_codes=scoped_control_codes,
                )
                sql_queries.append(
                    {
                        "label": "Comparison baseline query",
                        "sql": _render_sql_with_params(benchmark_sql, benchmark_params),
                    }
                )
                if benchmark_value is not None:
                    benchmark_scope = _format_control_scope(scoped_control_codes)
                    benchmark_label_suffix = benchmark_stat
                    if benchmark_scope and scoped_scope_label:
                        benchmark_label = f"{scoped_scope_label} {benchmark_scope.lower()} {benchmark_label_suffix}"
                    elif benchmark_scope:
                        benchmark_label = f"{benchmark_scope} {benchmark_label_suffix}"
                    else:
                        benchmark_label = f"{benchmark_label_suffix.capitalize()} {METRIC_DEFINITIONS[metric_key]['short_label'].lower()}"
                    benchmark = {
                        "label": benchmark_label,
                        "value": benchmark_value,
                        "scope": peer_scope,
                    }
                preferred_graph = "distribution" if peer_distribution else "comparison"
                auto_show_graph = bool(peer_distribution)
            elif result_kind == "school_profile":
                distribution_sql, distribution_params = _build_distribution_sql(
                    metric_key=metric_key,
                    state_filter=focus_row["state"] if focus_row.get("state") else None,
                    control_codes=(focus_row["control_code"],) if focus_row.get("control_code") else None,
                )
                peer_distribution, peer_scope = _query_peer_distribution(connection, focus_row, metric_key)
                sql_queries.append(
                    {
                        "label": "Graph distribution query",
                        "sql": _render_sql_with_params(distribution_sql, distribution_params),
                    }
                )
                preferred_graph = "distribution" if peer_distribution else "comparison"
                auto_show_graph = bool(peer_distribution)
            else:
                peer_distribution = []
                peer_scope = ""
                auto_show_graph = True
        else:
            ranking_sql, ranking_params = _build_ranked_institutions_sql(
                metric_key=metric_key,
                limit=limit,
                sort_direction=sort_direction,
                state_filter=state_filter,
                control_filter=control_filter,
            )
            rows = _query_ranked_institutions(
                connection,
                metric_key=metric_key,
                limit=limit,
                sort_direction=sort_direction,
                state_filter=state_filter,
                control_filter=control_filter,
            )
            if not rows:
                return None
            sql_queries.append(
                {
                    "label": "Primary result query",
                    "sql": _render_sql_with_params(ranking_sql, ranking_params),
                }
            )
            result_kind = "ranking"
            focus_row = rows[0]
            peer_distribution = []
            peer_scope = ""
            preferred_graph = "comparison"
            auto_show_graph = True

        if result_kind == "school_profile" and focus_row:
            net_price_history, history_lookup_key = _fetch_supplemental_net_price_history(
                connection,
                focus_row["institution_name"],
                limit_years=SUPPLEMENTAL_NET_PRICE_CHART_YEARS,
            )
            if net_price_history and history_lookup_key:
                history_sql = f"""
                    SELECT
                        MIN(school_name) AS school_name,
                        year_label,
                        year_start,
                        AVG(average_net_price) AS average_net_price,
                        AVG(net_price_income_0_30000) AS net_price_income_0_30000,
                        AVG(net_price_income_30001_48000) AS net_price_income_30001_48000,
                        AVG(net_price_income_48001_75000) AS net_price_income_48001_75000,
                        AVG(net_price_income_75001_110000) AS net_price_income_75001_110000,
                        AVG(net_price_income_110001_plus) AS net_price_income_110001_plus
                    FROM {SUPPLEMENTAL_NET_PRICE_TABLE}
                    WHERE normalized_name = ?
                    GROUP BY year_label, year_start
                    ORDER BY year_start DESC
                    LIMIT ?
                """
                sql_queries.append(
                    {
                        "label": "Supplemental net price history query",
                        "sql": _render_sql_with_params(
                            history_sql,
                            [history_lookup_key, SUPPLEMENTAL_NET_PRICE_CHART_YEARS],
                        ),
                    }
                )

    payload = {
        "question": question,
        "payload_id": _hash_value(
            json.dumps(
                {
                    "question": question,
                    "metric_key": metric_key,
                    "matched_schools": matched_schools,
                    "state_filter": state_filter,
                    "control_filter": control_filter,
                    "result_kind": result_kind,
                },
                sort_keys=True,
            )
        )[:12],
        "result_kind": result_kind,
        "metric_key": metric_key,
        "matched_schools": matched_schools,
        "state_filter": state_filter,
        "control_filter": control_filter,
        "assumptions": assumptions,
        "rows": rows,
        "focus_row": focus_row,
        "peer_distribution": peer_distribution,
        "peer_scope": peer_scope,
        "summary_text": _build_summary_text(
            result_kind=result_kind,
            metric_key=metric_key,
            rows=rows,
            benchmark=benchmark,
            requested_metrics=requested_metrics,
        ),
        "summary_stats": _build_summary_stats(
            result_kind,
            rows,
            metric_key=metric_key,
            benchmark=benchmark,
        ),
        "display_table": _build_display_table(metric_key, rows),
        "net_price_history": net_price_history,
        "auto_show_graph": auto_show_graph,
        "preferred_graph": preferred_graph,
        "benchmark": benchmark,
        "sql_queries": sql_queries,
        "requested_metrics": requested_metrics,
        "benchmark_stat": benchmark_stat,
    }
    payload["structured_context"] = _build_structured_context_from_payload(payload)
    return payload


def _build_primary_bar_chart(payload: dict, height: int) -> alt.Chart | None:
    metric_key = payload["metric_key"]
    metric_label = METRIC_DEFINITIONS[metric_key]["label"]
    chart_rows = [
        {
            "institution_name": row["institution_name"],
            "metric_value": row["metric_value"],
            "control": row["control"],
            "location": row["location"],
        }
        for row in payload["rows"]
        if row["metric_value"] is not None
    ]
    if not chart_rows:
        return None

    frame = pd.DataFrame(chart_rows).sort_values("metric_value", ascending=True)
    return (
        alt.Chart(frame)
        .mark_bar(color="#10523e", cornerRadiusEnd=4)
        .encode(
            x=alt.X("metric_value:Q", title=metric_label, axis=alt.Axis(format=_metric_axis_format(metric_key))),
            y=alt.Y("institution_name:N", sort=list(frame["institution_name"]), title=None),
            tooltip=[
                alt.Tooltip("institution_name:N", title="School"),
                alt.Tooltip("location:N", title="Location"),
                alt.Tooltip("control:N", title="Control"),
                alt.Tooltip("metric_value:Q", title=metric_label, format=_metric_axis_format(metric_key)),
            ],
        )
        .properties(height=max(height, 80 + 40 * len(frame)))
        .configure_view(strokeOpacity=0)
    )


def _build_range_chart(
    distribution: list[dict],
    metric_key: str,
    title: str,
    subtitle: str,
    height: int,
    benchmark: dict | None = None,
) -> alt.Chart | None:
    if not distribution:
        return None

    frame = pd.DataFrame(distribution)
    min_value = float(frame["metric_value"].min())
    max_value = float(frame["metric_value"].max())
    frame["lane"] = "Range"
    range_frame = pd.DataFrame([{"lane": "Range", "min_value": min_value, "max_value": max_value}])

    base = alt.Chart(frame).encode(
        x=alt.X(
            "metric_value:Q",
            title=METRIC_DEFINITIONS[metric_key]["label"],
            axis=alt.Axis(format=_metric_axis_format(metric_key)),
        ),
        y=alt.Y("lane:N", axis=None),
        tooltip=[
            alt.Tooltip("label:N", title="School"),
            alt.Tooltip(
                "metric_value:Q",
                title=METRIC_DEFINITIONS[metric_key]["label"],
                format=_metric_axis_format(metric_key),
            ),
        ],
    )

    range_rule = alt.Chart(range_frame).mark_rule(
        color="#c8c1ac",
        strokeWidth=6,
        opacity=0.65,
    ).encode(
        x="min_value:Q",
        x2="max_value:Q",
        y="lane:N",
    )

    comparison_ticks = base.transform_filter(
        alt.datum.is_selected == False
    ).mark_tick(color="#8cb885", thickness=2, size=30, opacity=0.85)

    selected_points = base.transform_filter(
        alt.datum.is_selected == True
    ).mark_point(color="#10523e", filled=True, size=170)

    chart_layers = [range_rule, comparison_ticks, selected_points]

    benchmark_value = benchmark.get("value") if benchmark else None
    if benchmark_value is not None:
        benchmark_frame = pd.DataFrame(
            [
                {
                    "lane": "Range",
                    "metric_value": benchmark_value,
                    "label": benchmark.get("label", "Benchmark"),
                }
            ]
        )
        benchmark_rule = (
            alt.Chart(benchmark_frame)
            .mark_rule(color="#f0a43a", strokeDash=[8, 4], strokeWidth=2.5)
            .encode(
                x="metric_value:Q",
                y="lane:N",
                tooltip=[
                    alt.Tooltip("label:N", title="Benchmark"),
                    alt.Tooltip(
                        "metric_value:Q",
                        title=METRIC_DEFINITIONS[metric_key]["label"],
                        format=_metric_axis_format(metric_key),
                    ),
                ],
            )
        )
        benchmark_point = (
            alt.Chart(benchmark_frame)
            .mark_point(color="#f0a43a", filled=True, size=120, shape="diamond")
            .encode(x="metric_value:Q", y="lane:N")
        )
        chart_layers.extend([benchmark_rule, benchmark_point])

    return (
        alt.layer(*chart_layers)
        .properties(height=height, title=alt.TitleParams(title, subtitle=subtitle))
        .configure_view(strokeOpacity=0)
    )


def _build_affordability_scatter(rows: list[dict], height: int) -> alt.Chart | None:
    scatter_rows = [
        {
            "institution_name": row["institution_name"],
            "net_price": row["net_price"],
            "earnings": row["earnings"],
            "undergrad_size": row["undergrad_size"],
            "control": row["control"],
            "location": row["location"],
        }
        for row in rows
        if row["net_price"] is not None and row["earnings"] is not None
    ]
    if len(scatter_rows) < 2:
        return None

    frame = pd.DataFrame(scatter_rows)
    return (
        alt.Chart(frame)
        .mark_circle(size=180, opacity=0.8)
        .encode(
            x=alt.X("net_price:Q", title=METRIC_DEFINITIONS["net_price"]["label"], axis=alt.Axis(format="$,.0f")),
            y=alt.Y("earnings:Q", title="Median Earnings (10 Years After Entry)", axis=alt.Axis(format="$,.0f")),
            color=alt.Color("control:N", title="Control"),
            size=alt.Size("undergrad_size:Q", title="Undergraduates"),
            tooltip=[
                alt.Tooltip("institution_name:N", title="School"),
                alt.Tooltip("location:N", title="Location"),
                alt.Tooltip("control:N", title="Control"),
                alt.Tooltip("net_price:Q", title=METRIC_DEFINITIONS["net_price"]["label"], format="$,.0f"),
                alt.Tooltip("earnings:Q", title="Earnings", format="$,.0f"),
                alt.Tooltip("undergrad_size:Q", title="Undergraduates", format=",.0f"),
            ],
        )
        .properties(height=height)
        .configure_view(strokeOpacity=0)
    )


def _build_result_chart(payload: dict, height: int) -> alt.Chart | None:
    if payload.get("preferred_graph") == "distribution" and payload.get("peer_distribution"):
        return _build_range_chart(
            payload["peer_distribution"],
            metric_key=payload["metric_key"],
            title=f"{METRIC_DEFINITIONS[payload['metric_key']]['label']} distribution",
            subtitle=f"{payload['focus_row']['institution_name']} highlighted within {payload['peer_scope']}",
            height=150,
            benchmark=payload.get("benchmark"),
        )

    if payload["result_kind"] == "school_profile":
        with sqlite3.connect(SQLITE_PATH) as connection:
            connection.row_factory = sqlite3.Row
            chart = _build_peer_affordability_scatter(
                connection,
                focus_school=payload["focus_row"],
                height=height,
            )
        if chart is not None:
            return chart
        return _build_primary_bar_chart(payload, height)

    if payload["metric_key"] in {"earnings", "net_price"}:
        chart = _build_affordability_scatter(payload["rows"], height)
        if chart is not None:
            return chart

    return _build_primary_bar_chart(payload, height)


def _build_peer_affordability_scatter(
    connection: sqlite3.Connection,
    focus_school: dict,
    height: int,
) -> alt.Chart | None:
    where_clauses = [
        "INSTNM IS NOT NULL",
        f"{NET_PRICE_SQL} IS NOT NULL",
        "MD_EARN_WNE_P10 IS NOT NULL",
        "UGDS IS NOT NULL",
    ]
    params: list[object] = []

    if focus_school["state"]:
        where_clauses.append("STABBR = ?")
        params.append(focus_school["state"])
    if focus_school["control_code"]:
        where_clauses.append("CONTROL = ?")
        params.append(focus_school["control_code"])

    rows = connection.execute(
        f"""
        SELECT
            INSTNM,
            CITY,
            STABBR,
            CONTROL,
            {NET_PRICE_SQL} AS net_price,
            MD_EARN_WNE_P10 AS earnings,
            UGDS AS undergrad_size
        FROM {SCORECARD_TABLE}
        WHERE {' AND '.join(where_clauses)}
        ORDER BY INSTNM
        """
        ,
        params,
    ).fetchall()

    normalized_focus = _normalize_lookup_value(focus_school["institution_name"])
    scatter_rows = []
    for row in rows:
        scatter_rows.append(
            {
                "institution_name": row["INSTNM"],
                "net_price": _coerce_float(row["net_price"]),
                "earnings": _coerce_float(row["earnings"]),
                "undergrad_size": _coerce_float(row["undergrad_size"]),
                "control": CONTROL_LABELS.get(str(row["CONTROL"]).strip(), "Unknown")
                if row["CONTROL"] is not None
                else "Unknown",
                "location": _format_location(row["CITY"], row["STABBR"]),
                "is_selected": _normalize_lookup_value(row["INSTNM"]) == normalized_focus,
            }
        )

    if len(scatter_rows) < 2:
        return None

    frame = pd.DataFrame(scatter_rows)
    base = alt.Chart(frame).encode(
        x=alt.X("net_price:Q", title=METRIC_DEFINITIONS["net_price"]["label"], axis=alt.Axis(format="$,.0f")),
        y=alt.Y(
            "earnings:Q",
            title="Median Earnings (10 Years After Entry)",
            axis=alt.Axis(format="$,.0f"),
        ),
        size=alt.Size("undergrad_size:Q", title="Undergraduates"),
        tooltip=[
            alt.Tooltip("institution_name:N", title="School"),
            alt.Tooltip("location:N", title="Location"),
            alt.Tooltip("control:N", title="Type of University"),
            alt.Tooltip("net_price:Q", title=METRIC_DEFINITIONS["net_price"]["label"], format="$,.0f"),
            alt.Tooltip("earnings:Q", title="Median earnings", format="$,.0f"),
            alt.Tooltip("undergrad_size:Q", title="Undergraduates", format=",.0f"),
        ],
    )

    comparison_points = base.transform_filter(
        alt.datum.is_selected == False
    ).mark_circle(color="#8cb885", opacity=0.55)
    selected_point = base.transform_filter(
        alt.datum.is_selected == True
    ).mark_circle(color="#10523e", opacity=0.95, stroke="white", strokeWidth=1.5)

    return (
        (comparison_points + selected_point)
        .properties(height=height)
        .configure_view(strokeOpacity=0)
    )


def _build_supplemental_net_price_history_chart(
    history_rows: list[dict],
    view_key: str,
    height: int = 240,
) -> alt.Chart | None:
    view_label = SUPPLEMENTAL_NET_PRICE_VIEW_LABELS.get(view_key, "Average net price")
    chart_rows = [
        {
            "year_label": row["year_label"],
            "year_start": row["year_start"],
            "metric_value": row.get(view_key),
        }
        for row in history_rows
        if row.get(view_key) is not None
    ]
    if not chart_rows:
        return None

    frame = pd.DataFrame(chart_rows).sort_values("year_start")
    x_order = frame["year_label"].tolist()
    line = (
        alt.Chart(frame)
        .mark_line(color="#10523e", strokeWidth=3)
        .encode(
            x=alt.X("year_label:O", title="Academic year", sort=x_order),
            y=alt.Y("metric_value:Q", title=view_label, axis=alt.Axis(format="$,.0f")),
            tooltip=[
                alt.Tooltip("year_label:O", title="Academic year"),
                alt.Tooltip("metric_value:Q", title=view_label, format="$,.0f"),
            ],
        )
    )
    points = alt.Chart(frame).mark_point(color="#10523e", filled=True, size=95).encode(
        x=alt.X("year_label:O", sort=x_order),
        y="metric_value:Q",
        tooltip=[
            alt.Tooltip("year_label:O", title="Academic year"),
            alt.Tooltip("metric_value:Q", title=view_label, format="$,.0f"),
        ],
    )
    return (line + points).properties(height=height).configure_view(strokeOpacity=0)


def _render_supplemental_net_price_history_panel(payload: dict) -> None:
    if payload.get("result_kind") != "school_profile":
        return

    history_rows = payload.get("net_price_history") or []
    if not history_rows:
        return

    payload_id = payload.get("payload_id") or _hash_value(payload.get("question", ""))[:12]
    selected_views = st.session_state.setdefault("net_price_history_view_by_payload", {})
    selected_view = selected_views.get(payload_id, SUPPLEMENTAL_NET_PRICE_DEFAULT_VIEW)

    selected_view = st.pills(
        "View Average Net Price by Family Income",
        options=[view_key for view_key, _ in SUPPLEMENTAL_NET_PRICE_VIEW_OPTIONS],
        default=selected_view,
        format_func=lambda view_key: SUPPLEMENTAL_NET_PRICE_VIEW_LABELS.get(view_key, str(view_key)),
        key=f"supplemental_net_price_view_{payload_id}",
        selection_mode="single",
        width="content",
    ) or SUPPLEMENTAL_NET_PRICE_DEFAULT_VIEW
    selected_views[payload_id] = selected_view

    chart = _build_supplemental_net_price_history_chart(history_rows, selected_view)
    if chart is not None:
        st.altair_chart(chart, use_container_width=True)
    else:
        st.info("This school does not have supplemental history data for that income bracket.")

    focus_row = payload.get("focus_row") or {}
    if focus_row.get("net_price_source") == "supplemental_history":
        st.caption(
            "The average net price shown above comes from U.S. FAFSA data."
        )
    else:
        st.caption(
            "Historical net price values shown here come from the supplemental in-state tuition history source."
        )


def _render_sql_result_panel(payload: dict, show_header: bool = True) -> None:
    if show_header:
        st.subheader("Data Results")

    summary_stats = payload.get("summary_stats", [])
    if summary_stats:
        columns = st.columns(len(summary_stats))
        for column, stat in zip(columns, summary_stats):
            column.markdown(
                f"""
                <div class="edu-summary-stat">
                    <div class="edu-summary-label">{stat['label']}</div>
                    <div class="edu-summary-value">{stat['value']}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )

    _render_supplemental_net_price_history_panel(payload)
    st.dataframe(payload["display_table"], use_container_width=True, hide_index=True)

    payload_id = payload.get("payload_id") or _hash_value(payload.get("question", ""))[:12]
    visible_chart_ids = st.session_state.setdefault("visible_chart_ids", {})
    should_show_chart = visible_chart_ids.get(payload_id, payload.get("auto_show_graph", False))
    button_label = "Hide comparison graph" if should_show_chart else "View comparison graph"
    if st.button(button_label, key=f"toggle_chart_{payload_id}"):
        visible_chart_ids[payload_id] = not should_show_chart
        st.rerun()

    if should_show_chart:
        chart = _build_result_chart(payload, 320)

        if chart is not None:
            st.altair_chart(chart, use_container_width=True)
        else:
            st.info("There was not enough data to build a chart for this result.")

    query_items = _build_sql_query_display_items(payload)
    expander_label = "SQL Query Used" if len(query_items) == 1 else "SQL Queries Used"
    with st.expander(expander_label, expanded=False):
        if query_items:
            for index, item in enumerate(query_items):
                if len(query_items) > 1:
                    st.caption(item["label"])
                st.code(item["sql"], language="sql")
                if len(query_items) > 1 and index < len(query_items) - 1:
                    st.divider()
        else:
            st.code(payload["structured_context"])

    if payload.get("summary_text"):
        st.markdown(
            _format_text_block_html(payload["summary_text"], "edu-result-summary"),
            unsafe_allow_html=True,
        )


def _render_conversation_history() -> None:
    history = st.session_state.get("chat_history", [])
    if not history:
        return

    for entry in history:
        with st.chat_message(entry["role"]):
            if entry["role"] == "user":
                st.write(entry["content"])
                continue

            if entry.get("kind") == "sql":
                _render_sql_result_panel(entry["payload"], show_header=True)
            else:
                _render_assistant_text(entry["content"])
                _render_rag_watermark(entry.get("rag_citations"))
                if entry.get("structured_context"):
                    with st.expander("Structured SQLite Context Used", expanded=False):
                        st.code(entry["structured_context"])


def _render_session_summary(summary: dict | None) -> None:
    if not summary:
        return

    takeaways = summary.get("takeaways") or summary.get("findings") or []
    next_steps = summary.get("next_steps") or summary.get("follow_ups") or []
    closing_thought = str(summary.get("closing_thought", "")).strip()

    if not takeaways:
        return

    st.subheader("Session Wrap-Up")
    with st.container(border=False):
        st.markdown("**Top Takeaways**")
        for item in takeaways[:3]:
            st.write(f"- {item}")

        st.markdown("**Next Steps**")
        for item in next_steps[:3]:
            st.write(f"- {item}")

        if closing_thought:
            st.markdown("**Closing Thought**")
            st.write(closing_thought)


def _build_sqlite_context(
    question: str,
    context_schools: tuple[str, ...] | None = None,
    context_metric_key: str | None = None,
) -> str:
    payload = _build_sql_result_payload(
        question,
        context_schools=context_schools,
        context_metric_key=context_metric_key,
    )
    return payload["structured_context"] if payload else ""


def _build_application_information_context() -> str:
    field_map = [
        ("app_info_gpa", "GPA (out of 4)"),
        ("app_info_sat", "SAT"),
        ("app_info_act", "ACT"),
        ("app_info_home_state", "Home state"),
        ("app_info_major_interest", "Major interest"),
    ]

    lines = []
    for state_key, label in field_map:
        value = str(st.session_state.get(state_key, "")).strip()
        if value:
            lines.append(f"- {label}: {value}")

    if not lines:
        return ""

    return "Application information provided by the student:\n" + "\n".join(lines)

def _build_rag_payload(question: str) -> dict:
    empty_payload = {"context": "", "citations": []}
    if not looks_like_rag_question(question):
        return empty_payload

    try:
        payload = retrieve_rag_context(question)
    except Exception:
        logger.exception("RAG retrieval failed.")
        return empty_payload

    context = str(payload.get("context", "")).strip()
    matches = payload.get("matches", [])
    citations = []
    seen = set()

    for match in matches:
        cip_title = str(match.get("cip_title", "")).strip()
        degree_level = str(match.get("degree_level", "")).strip()
        cip4 = str(match.get("cip4", "")).strip()
        if not cip_title:
            continue

        citation = cip_title
        if degree_level:
            citation += f" ({degree_level}"
            if cip4:
                citation += f", CIP {cip4}"
            citation += ")"
        elif cip4:
            citation += f" (CIP {cip4})"

        if citation in seen:
            continue
        citations.append(citation)
        seen.add(citation)
        if len(citations) == 3:
            break

    return {
        "context": context,
        "citations": citations,
    }


def _render_rag_watermark(citations: list[str] | None) -> None:
    items = [str(item).strip() for item in (citations or []) if str(item).strip()]
    if not items:
        return
    escaped_items = [html.escape(item) for item in items]

    st.markdown(
        (
            "<div class='edu-rag-watermark'>"
            "<strong>Retrieved program context used.</strong> "
            f"Sources: {'; '.join(escaped_items)}"
            "</div>"
        ),
        unsafe_allow_html=True,
    )


def _build_model_routing_message(
    source_plan: dict,
    structured_context: str,
    rag_context: str,
    sql_fallback_instruction: str = "",
) -> str:
    active_sources = []
    if source_plan.get("use_resume"):
        active_sources.append("resume context")
    if source_plan.get("use_applicant_stats"):
        active_sources.append("applicant stats")
    if structured_context:
        active_sources.append("school-level SQLite context")
    if rag_context:
        active_sources.append("retrieved program context")
    if sql_fallback_instruction:
        active_sources.append("general reasoning with repository-gap fallback")
    if not active_sources:
        active_sources.append("general reasoning")

    lines = [
        f"Current prompt classification: {source_plan['intent']}.",
        "Active sources for this answer: " + ", ".join(active_sources) + ".",
        "Do not mention internal routing.",
    ]

    if source_plan["intent"] == INTENT_PERSONALIZED_ADVICE:
        if source_plan.get("use_resume") or source_plan.get("use_applicant_stats"):
            lines.append(
                "Tailor the answer to the available personal context, but keep the response grounded and low-pressure."
            )
        else:
            lines.append(
                "No additional stored personal context is currently available beyond the user's prompt. Give helpful general guidance and only optionally suggest sharing GPA, SAT/ACT, home state, major interest, or a resume if that would materially improve the answer."
            )

    if structured_context and source_plan.get("use_sql_context"):
        lines.append(
            "Use the school-level structured context as supporting evidence rather than turning the answer into a raw data dump."
        )
    if rag_context and source_plan.get("use_rag"):
        lines.append(
            "Use the retrieved program context for claims about majors, careers, salaries, occupations, or degree-level outcomes."
        )
    if sql_fallback_instruction:
        lines.append(sql_fallback_instruction)

    return " ".join(lines)


def _build_messages(
    user_input: str,
    resume_text: str,
    structured_context: str = "",
    rag_context: str = "",
    sql_fallback_instruction: str = "",
    source_plan: dict | None = None,
    chat_history: list[dict] | None = None,
) -> list[dict]:
    messages = [{"role": "system", "content": BASE_SYSTEM_MESSAGE}]
    if source_plan:
        messages.append(
            {
                "role": "system",
                "content": _build_model_routing_message(
                    source_plan,
                    structured_context=structured_context,
                    rag_context=rag_context,
                    sql_fallback_instruction=sql_fallback_instruction,
                ),
            }
        )
    if structured_context:
        messages.append(
            {
                "role": "system",
                "content": _truncate_text(
                    SQLITE_CONTEXT_SYSTEM_MESSAGE + "\n\n" + structured_context,
                    6000,
                ),
            }
        )
    if rag_context:
        messages.append(
            {
                "role": "system",
                "content": _truncate_text(
                    RAG_CONTEXT_SYSTEM_MESSAGE + "\n\n" + rag_context,
                    6000,
                ),
            }
        )
    messages.append(
        {
            "role": "system",
            "content": AZURE_RESPONSE_STYLE_SYSTEM_MESSAGE,
        }
    )
    application_info_context = _build_application_information_context()
    if source_plan and source_plan.get("use_applicant_stats") and application_info_context:
        messages.append(
            {
                "role": "system",
                "content": application_info_context,
            }
        )
    if source_plan and source_plan.get("use_resume") and resume_text:
        messages.append(
            {
                "role": "system",
                "content": "Personalize answers using this resume context:\n"
                f"{_truncate_text(resume_text, 3000)}",
            }
        )
    if chat_history:
        messages.extend(_history_to_model_messages(chat_history))
    messages.append({"role": "user", "content": user_input})
    return messages


def _render_resume_panel() -> None:
    with st.expander("Resume Context", expanded=bool(st.session_state.get("resume_text"))):
        uploaded_resume = st.file_uploader(
            "Upload a resume",
            type=["pdf", "docx", "txt"],
            help="Supported formats: PDF, DOCX, TXT.",
            key="resume_uploader",
        )
        if uploaded_resume is None:
            if st.session_state.get("resume_text") or st.session_state.get("resume_filename"):
                _clear_resume_state()
        elif uploaded_resume.name != st.session_state.get("resume_filename"):
            _process_uploaded_resume(uploaded_resume)

        if st.session_state.get("resume_text"):
            st.caption(f"Loaded: {st.session_state['resume_filename']}")
            st.caption("Preview below is redacted before storage and model use.")
            st.text_area(
                "Redacted preview",
                value=_truncate_text(st.session_state["resume_text"], 900),
                height=220,
                disabled=True,
            )
        else:
            st.caption("No resume uploaded.")


def _handle_pending_prompt() -> None:
    user_input = st.session_state.get("pending_prompt", "")
    if not user_input:
        return

    prior_history = st.session_state.get("chat_history", [])[:-1]
    prior_sql_payload = _get_recent_sql_payload(prior_history)
    current_matched_schools = _match_schools(user_input, _load_school_names())
    current_state_filter = _extract_state_filter(user_input, ignore_phrases=current_matched_schools)
    current_control_filter = _extract_control_filter(user_input)
    _, current_metric_detected, _ = _resolve_metric(user_input)
    current_wants_ranked_results = _wants_ranked_results(user_input)
    current_wants_distribution_graph = _wants_distribution_graph(user_input)

    context_schools = ()
    if _should_reuse_school_context(
        user_input,
        prior_sql_payload=prior_sql_payload,
        matched_schools=current_matched_schools,
        state_filter=current_state_filter,
        control_filter=current_control_filter,
    ):
        context_schools = tuple(_extract_context_schools_from_payload(prior_sql_payload))

    context_metric_key = None
    if _should_reuse_metric_context(
        user_input,
        prior_sql_payload=prior_sql_payload,
        metric_detected=current_metric_detected,
        matched_schools=current_matched_schools,
        state_filter=current_state_filter,
        control_filter=current_control_filter,
        wants_ranked_results=current_wants_ranked_results,
        wants_distribution_graph=current_wants_distribution_graph,
    ):
        prior_metric_key = prior_sql_payload.get("metric_key") if prior_sql_payload else None
        if prior_metric_key in METRIC_DEFINITIONS:
            context_metric_key = prior_metric_key

    source_plan = _build_source_plan(
        user_input,
        matched_schools=current_matched_schools,
        metric_detected=current_metric_detected,
        wants_ranked_results=current_wants_ranked_results,
        wants_distribution_graph=current_wants_distribution_graph,
        state_filter=current_state_filter,
        control_filter=current_control_filter,
        context_schools=context_schools,
    )

    with st.chat_message("assistant"):
        sql_fallback_instruction = ""
        if source_plan["use_sql_result"]:
            with st.spinner("Checking school data..."):
                sql_result = _build_sql_result_payload(
                    user_input,
                    context_schools=context_schools,
                    context_metric_key=context_metric_key,
                )
            if sql_result:
                _render_sql_result_panel(sql_result)
                st.session_state.last_sql_result = sql_result
                st.session_state.last_reply = ""
                st.session_state.last_structured_context = sql_result["structured_context"]
                _append_chat_history_entry(
                    {
                        "role": "assistant",
                        "kind": "sql",
                        "payload": sql_result,
                        "summary_text": sql_result["summary_text"],
                        "structured_context": sql_result["structured_context"],
                        "content": sql_result["summary_text"],
                    }
                )
                st.session_state.pending_prompt = ""
                return

            unmatched_school_message = _build_unmatched_school_message(user_input)
            no_data_message = _build_no_data_message(
                user_input,
                context_schools=context_schools,
                context_metric_key=context_metric_key,
            )
            if unmatched_school_message or no_data_message:
                sql_fallback_instruction = _build_sql_fallback_instruction(
                    unmatched_school_message=unmatched_school_message,
                    no_data_message=no_data_message,
                )

        st.session_state.last_sql_result = None
        with st.spinner("Preparing reply..."):
            structured_context = ""
            if source_plan["use_sql_context"]:
                structured_context = _build_sqlite_context(
                    user_input,
                    context_schools=context_schools,
                    context_metric_key=context_metric_key,
                )
            rag_payload = {"context": "", "citations": []}
            if source_plan["use_rag"]:
                rag_payload = _build_rag_payload(user_input)
            rag_context = rag_payload["context"]
            rag_citations = rag_payload["citations"]
        st.session_state.last_structured_context = structured_context

        openai.api_type = "azure"
        openai.api_base = os.getenv("AZURE_OPENAI_ENDPOINT")
        openai.api_version = os.getenv("AZURE_OPENAI_API_VERSION", "2025-04-01-preview")

        api_key = os.getenv("AZURE_OPENAI_API_KEY")
        deployment_id = os.getenv("AZURE_OPENAI_DEPLOYMENT_ID")
        if not api_key or not deployment_id:
            error_message = (
                "Set `AZURE_OPENAI_API_KEY` and `AZURE_OPENAI_DEPLOYMENT_ID` to generate non-SQL chat replies."
            )
            st.error(error_message)
            _append_chat_history_entry(
                {
                    "role": "assistant",
                    "kind": "model",
                    "content": error_message,
                    "structured_context": structured_context,
                }
            )
            st.session_state.pending_prompt = ""
            return

        openai.api_key = api_key
        messages = _build_messages(
            user_input,
            st.session_state.get("resume_text", ""),
            structured_context=structured_context,
            rag_context=rag_context,
            sql_fallback_instruction=sql_fallback_instruction,
            source_plan=source_plan,
            chat_history=prior_history,
        )
        reply = _render_streamed_assistant_reply(deployment_id, messages)
        _render_rag_watermark(rag_citations)

        st.session_state.last_reply = reply
        _append_chat_history_entry(
            {
                "role": "assistant",
                "kind": "model",
                "content": reply,
                "structured_context": structured_context,
                "rag_citations": rag_citations,
            }
        )
        st.session_state.pending_prompt = ""
        return


def main() -> None:
    st.session_state.setdefault("resume_text", "")
    st.session_state.setdefault("resume_filename", "")
    st.session_state.setdefault("app_info_gpa", "")
    st.session_state.setdefault("app_info_sat", "")
    st.session_state.setdefault("app_info_act", "")
    st.session_state.setdefault("app_info_home_state", "")
    st.session_state.setdefault("app_info_major_interest", "")
    st.session_state.setdefault("last_reply", "")
    st.session_state.setdefault("last_structured_context", "")
    st.session_state.setdefault("last_sql_result", None)
    st.session_state.setdefault("chat_history", [])
    st.session_state.setdefault("session_summary", None)
    st.session_state.setdefault("pending_prompt", "")
    st.session_state.setdefault("visible_chart_ids", {})
    st.session_state.setdefault("net_price_history_view_by_payload", {})

    _apply_brand_theme()
    end_session_clicked = _render_sidebar_panel()

    if end_session_clicked:
        if not st.session_state.get("chat_history"):
            st.warning("There is no active chat history to summarize yet.")
            return

        with st.spinner("Wrapping up the session..."):
            summary = _generate_session_summary(st.session_state["chat_history"])

        st.session_state.session_summary = summary
        st.session_state.chat_history = []
        st.session_state.last_reply = ""
        st.session_state.last_structured_context = ""
        st.session_state.last_sql_result = None
        st.session_state.pending_prompt = ""
        st.session_state.visible_chart_ids = {}
        st.session_state.net_price_history_view_by_payload = {}
        st.rerun()

    conversation_placeholder = st.container()
    prompt = st.chat_input(
        "Ask about tuition, net price, debt, admission rate, completion, or long-term earnings."
    )
    if prompt and prompt.strip():
        st.session_state.session_summary = None
        st.session_state.pending_prompt = prompt.strip()
        _append_chat_history_entry({"role": "user", "content": prompt.strip()})

    with conversation_placeholder:
        st.markdown("<div class='edu-chat-top-guard'></div>", unsafe_allow_html=True)
        _render_session_summary(st.session_state.get("session_summary"))
        _render_conversation_history()
        if st.session_state.get("pending_prompt"):
            _handle_pending_prompt()


if __name__ == "__main__":
    main()
